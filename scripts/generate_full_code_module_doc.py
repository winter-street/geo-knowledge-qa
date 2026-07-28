"""Generate the full source-code module guide for the Geo-Knowledge Q&A project."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from generate_teacher_module_doc import (
    add_bullet,
    add_module,
    add_number,
    add_table,
    set_run_font,
    setup_doc,
)


OUT = Path(__file__).resolve().parents[1] / "docs" / "地质找矿智能问答系统-全量代码模块说明（详细版）.docx"


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(68)
    r = p.add_run("地质找矿智能问答系统")
    set_run_font(r, size=25, bold=True, color="1E3A5F")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("全量代码模块说明（详细版）")
    set_run_font(r, size=18, bold=True, color="2E7D5B")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("按当前运行源码逐文件整理：职责、输入、输出、调用关系与运行位置")
    set_run_font(r, size=10, color="6B6B78")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(112)
    r = p.add_run("2026 年 7 月")
    set_run_font(r, size=10, color="6B6B78")
    doc.add_page_break()


def add_file_catalog(doc, title, rows):
    doc.add_heading(title, level=2)
    doc.add_paragraph("表中“运行位置”用于区分：在线必经、离线建库、训练/验证、页面展示或仅作配置。")
    add_table(doc, ["文件", "职责", "主要输入", "主要输出", "运行位置"], rows, [4.1, 5.1, 3.2, 3.4, 2.0])


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_doc(doc)
    add_cover(doc)

    doc.add_heading("阅读范围与使用方法", level=1)
    doc.add_paragraph(
        "这份文档面向需要解释“每个功能是做什么、代码怎样跑”的场景。它覆盖当前项目中实际参与运行、"
        "建库、模型训练和验证的源码文件。每个核心文件都说明了输入、内部处理、输出和调用方；"
        "工具、测试和配置文件采用逐文件目录式说明。"
    )
    doc.add_heading("纳入范围", level=2)
    add_bullet(doc, "前端：frontend/src、Vite 配置、前端工具函数及现有 TypeScript 断言测试。")
    add_bullet(doc, "后端：backend/src、后端配置、问答调试脚本。")
    add_bullet(doc, "数据端：Flask 服务、PDF 到知识图谱的离线管线、NER 训练与标注脚本、测试与环境检查脚本。")
    add_bullet(doc, "根目录启动入口和各端依赖声明。")
    doc.add_heading("不纳入逐行解释的内容", level=2)
    add_bullet(doc, "node_modules、Python 虚拟环境、模型权重、SQLite 数据库、Neo4j dump、PDF 原件和其他生成产物。这些是依赖或数据，不是手写业务源码。")
    add_bullet(doc, "frontend/a的一些补充 下的旧前端备份与 design-demos。它们不在当前运行入口的调用链上。")
    doc.add_paragraph("注意：本文说明的是当前代码实际行为，不把“计划接入”写成“已经在在线问答中调用”。", style="Intense Quote")

    doc.add_heading("一、总运行图：代码怎样串起来", level=1)
    doc.add_paragraph("系统必须区分两条链路。离线链路先把资料准备好；在线链路只读取这些已经准备好的数据，不会每次提问都重新解析 PDF 或重新训练模型。")
    add_table(doc, ["链路", "入口代码", "主要调用", "最终产物"], [
        ["离线建库", "ml-service/scripts/run_pipeline.py", "parse_pdf -> chunk_text -> build_index -> embed_chunks -> extract_hybrid -> write_neo4j -> build_ontology -> run_reasoning", "SQLite 文档库、BGE 向量、Neo4j 图谱、OWL 本体、推理关系"],
        ["在线问答", "frontend QaView/store -> backend routes/qa.ts", "Flask /search + Neo4j searchEntities + LLM gateway -> SSE", "流式回答、来源、图谱路径、空间数据、问答日志"],
        ["后台管理", "frontend AdminView -> backend admin/upload routes", "SQLite 文档/日志、Neo4j 实体、后台 Python 管线", "统计、资料清单、上传结果、实体维护"],
    ], [3.0, 4.2, 7.7, 5.0])
    doc.add_heading("1.1 在线问答的固定调用顺序", level=2)
    for text in [
        "ChatInput.vue 收集问题和检索模式，QaView 调用 qa store 的 sendMessageStream。",
        "api/qa.ts 带 JWT 发起 POST /api/qa/ask，使用 fetch 读取 SSE。",
        "backend middleware/auth.ts 验证 JWT；routes/qa.ts 同时启动文献检索、图谱检索和空间检索。",
        "services/tfidf.ts 调 Flask /search；Flask 从 SQLite 的 BGE 或 TF-IDF 索引返回相关 chunks。",
        "services/kg.ts 直接查 Neo4j，返回关系路径与可能存在的坐标。",
        "services/llm.ts 和 gateway.ts 将两类证据写入提示词，调用大模型或使用检索原文降级。",
        "routes/qa.ts 依次回传 meta、chunk、done；store 把文字追加到 AI 气泡，ChatMessage/QaView 渲染来源、关系和地图。",
    ]:
        add_number(doc, text)
    doc.add_heading("1.2 数据端的固定调用顺序", level=2)
    for text in [
        "parse_pdf.py 产生带页码的纯文本。",
        "chunk_text.py 产生带页码和字符范围的重叠文本块。",
        "build_index.py 写入 documents/chunks 并训练 TF-IDF；embed_chunks.py 以 BGE 512 维向量覆盖 chunks.vector。",
        "extract_hybrid.py 用词典正则识别主体实体，BERT-NER 对词典覆盖较弱页面补漏，LLM 对已发现实体抽取关系。",
        "write_neo4j.py 将文档、实体和关系写入 Neo4j。",
        "build_ontology.py 生成 OWL；run_reasoning.py 推导关系，并把 inferred=true 的结果写回 Neo4j。",
    ]:
        add_number(doc, text)

    doc.add_heading("二、前端源码：页面、交互与状态", level=1)
    doc.add_paragraph("前端以 Vue 3 单页应用实现。这里的“输入”主要是用户动作、后端 API 数据和本地缓存；“输出”是 DOM 页面或发给后端的 HTTP 请求。")
    add_file_catalog(doc, "2.1 前端入口、类型、样式与构建配置", [
        ["frontend/src/main.ts", "应用启动入口", "App、Pinia、Router", "挂载到 #app 的 Vue 应用", "浏览器启动一次"],
        ["frontend/src/App.vue", "根组件", "router-view", "当前路由对应页面", "所有页面共同入口"],
        ["frontend/src/types/index.ts", "前端共享类型契约", "无运行输入", "Message、Source、SpatialData、Document、QaAskResponse 等类型", "编译期"],
        ["frontend/src/styles/global.css", "全局设计 Token 和 Element Plus 覆盖", "CSS 变量", "统一字体、颜色、边框和组件基础样式", "所有页面"],
        ["frontend/vite.config.ts", "Vite 开发与打包配置", "环境、插件配置", "@ 别名、自动导入、/api -> :3000 代理", "开发/构建"],
        ["frontend/package.json", "前端依赖和命令定义", "npm 命令", "dev、build、type-check、preview", "开发/构建"],
        ["frontend/DESIGN.md", "视觉规范说明", "无", "设计 Token 和反模式约束", "开发参考"],
    ])
    add_module(doc, "2.2 路由与双角色布局", "按照 user 与 admin 角色组织页面，并阻止越权访问。", "浏览器地址、localStorage 中 token/role。", "router/index.ts 用 hash 路由定义 /user 和 /admin 两套嵌套路由。beforeEach 先确认登录，再确认角色；访问登录页的已登录用户被送到各自首页。UserLayout 使用顶部导航；AdminLayout 使用可收起的侧边栏和顶栏。", "正确的页面组件或重定向。", "普通用户没有管理控制台入口；管理员能访问资料、实体、日志等页面。", "frontend/src/router/index.ts；frontend/src/layouts/UserLayout.vue；frontend/src/layouts/AdminLayout.vue")
    add_module(doc, "2.3 用户状态：stores/user.ts", "集中保存登录身份，并提供登录和退出方法。", "登录接口返回的 token、username、role；退出操作。", "Pinia store 在初始化时从 localStorage 恢复身份。login 调用 api/auth.ts，成功后保存三项数据；logout 清除本地身份并由布局组件跳回登录页。", "响应式身份状态和持久化后的 localStorage。", "路由守卫、axios 拦截器和布局都能使用同一份身份信息。", "frontend/src/stores/user.ts；frontend/src/api/auth.ts")
    add_module(doc, "2.4 请求基础设施：api/request.ts", "为普通 JSON API 提供统一请求方式。", "调用方给出的路径、请求体；localStorage token。", "创建 axios 实例，baseURL 为 /api。请求拦截器自动写入 Bearer token；响应拦截器遇到 401 清理登录态并跳转登录页。", "后端 JSON 响应或 Promise 错误。", "除 SSE 外的接口不需要重复处理 token 与未登录错误。", "frontend/src/api/request.ts")
    add_module(doc, "2.5 问答 API：api/qa.ts", "支持普通问答和流式问答，是前端与核心问答接口的协议实现。", "question、retrievalMode，以及流式回调。", "askQuestion 用 axios 取完整 JSON。askQuestionStream 用 fetch 发送 stream:true，逐行解析 SSE 的 data JSON；meta 转给 onMeta，chunk 转给 onChunk，done 携带 sources/kgContext/spatialData。normalizeSpatial 过滤不含合法 lng/lat 或折线路径的数据。", "完整 QaAskResponse，或取消函数和持续回调。", "后端可以边生成边推送，页面无须等整段答案生成完成。", "frontend/src/api/qa.ts")
    add_module(doc, "2.6 问答状态：stores/qa.ts", "维护会话、消息、检索模式和流式更新。", "用户文本、当前会话 ID、SSE 事件。", "create/select/delete/rename/pin 管理会话；sendMessage 处理非流式结果；sendMessageStream 先插入用户消息和空 AI 消息，再用 patchAiMsg 替换消息对象触发 Vue 响应式更新。完成后以 conversations 与 messagesMap 保存到 localStorage。", "currentMessages、loading、error、持久化的会话状态。", "对话切换不会丢失消息，AI 文字能逐段显示。", "frontend/src/stores/qa.ts")
    add_file_catalog(doc, "2.7 前端 API 文件逐项说明", [
        ["api/auth.ts", "登录请求封装", "username,password", "LoginResponse", "登录页 -> 后端 /auth/login"],
        ["api/doc.ts", "文档列表和上传请求", "FormData/PDF", "Document 列表或上传结果", "管理员控制台"],
        ["api/kg.ts", "子图和空间数据请求", "keyword, depth", "nodes/edges 或 SpatialData", "图谱、地图页"],
        ["api/qa.ts", "普通与 SSE 问答", "question, mode, callbacks", "回答/事件/取消函数", "问答页"],
        ["api/request.ts", "Axios 实例与拦截器", "任意 JSON 请求", "统一 API Promise", "所有非 SSE API"],
    ])
    add_module(doc, "2.8 登录页：LoginView.vue", "提供普通用户和管理员的演示登录入口。", "角色选择、用户名、密码。", "selectedRole 切换时填入相应演示账号；handleLogin 调 user store 登录，成功后按 role 跳转到 /user 或 /admin。", "登录成功后的首页，或页面错误提示。", "演示时可以快速切换两个身份验证路由和权限。", "frontend/src/views/LoginView.vue")
    add_module(doc, "2.9 问答页面：QaView.vue", "组装对话列表、聊天区、输入组件和问答地图。", "qa store 的会话/消息；发送、会话右键菜单、地图查看动作。", "处理会话置顶、改名、删除和筛选。监听消息中的 spatialData，按需加载高德 JS API，在点、折线和图层间维护覆盖物；无密钥或加载失败时使用文字兜底。", "完整问答界面及地图标注。", "问答的证据、关系和空间信息在同一工作区可见。", "frontend/src/views/QaView.vue")
    add_module(doc, "2.10 图谱页面：GraphView.vue", "展示关键词周边的 Neo4j 子图。", "关键词、跳数、后端返回的 nodes/edges、节点点击。", "loadGraph 请求 /kg/subgraph。renderGraph 找中心种子节点，以 computeHopDistances 做 BFS 计算跳数，为不同跳数分配同心圆初始位置，再用 G6 d3-force 布局渲染。节点点击显示相关边。", "交互式 G6 图、实体/关系统计和节点详情。", "用户可从“关系路径”理解为“关系网络”。", "frontend/src/views/GraphView.vue")
    add_module(doc, "2.11 地图页面：MapView.vue", "提供地图底图、绘制和空间查询入口。", "高德地图密钥、工具栏动作、关键词、SpatialData。", "初始化地图和控件；searchSpatial 调 KG 空间接口，renderSpatialData 绘制点/线；绘制勘查区、切换卫星底图和清除覆盖物由相应工具函数处理。未接入的 GeoJSON 空间分析功能明确给出未就绪提示。", "地图覆盖物、图层状态和提示。", "可演示空间展示；同时不把尚未接入的空间查询包装成已完成能力。", "frontend/src/views/MapView.vue")
    add_module(doc, "2.12 管理台：AdminView.vue", "让管理员查看知识库、日志和实体，并触发 PDF 上传。", "管理 API 的统计/文档/实体/日志；用户选择的 PDF；编辑后的实体字段。", "加载统计与列表；查看文档时请求 chunks；查看、修改、删除实体；上传用 FormData 调 uploadDoc。页面把状态、满意度、服务配置映射为标签和图表数据。", "统计卡片、表格、详情抽屉/弹窗和上传状态。", "管理员不必直接操作 SQLite 或 Neo4j 就能了解系统。", "frontend/src/views/AdminView.vue")
    add_file_catalog(doc, "2.13 可复用组件逐项说明", [
        ["components/ChatInput.vue", "问答输入框和模式选择", "文本、Enter/发送、rag/kg/hybrid", "send 事件和 selected mode", "QaView"],
        ["components/ChatMessage.vue", "单条消息展示", "Message、是否最后一条", "Markdown、来源、关系和地图查看事件", "QaView 消息列表"],
        ["components/KgMiniGraph.vue", "回答内嵌的小型图谱", "KGPath、问题、自动展开参数", "G6 小图/全屏图", "ChatMessage"],
        ["components/SourceCard.vue", "来源片段卡片", "Source", "标题、页码、摘要", "ChatMessage"],
        ["components/StatCard.vue", "统计数卡片", "label/value/unit", "统一统计展示", "AdminView"],
        ["components/ChartPanel.vue", "ECharts 容器", "图表 option", "可自适应图表", "AdminView"],
    ])
    add_file_catalog(doc, "2.14 前端工具与测试逐项说明", [
        ["utils/format.ts", "时间格式化", "时间戳", "适合界面展示的时间文本", "页面工具"],
        ["utils/markdown.ts", "安全的简化 Markdown 渲染", "回答文本", "HTML 字符串", "ChatMessage"],
        ["utils/markdown.test.ts", "Markdown 断言", "测试样例", "通过/失败信息", "手动 tsx 测试"],
        ["utils/search.ts", "会话标题过滤", "会话数组、关键词", "匹配后的会话数组", "问答会话筛选"],
        ["utils/search.test.ts", "会话搜索断言", "测试数据", "通过/失败信息", "手动 tsx 测试"],
        ["utils/spatial.ts", "空间类别元数据与示例数据", "实体类型、SpatialData", "颜色、名称、预设空间数据", "地图和图谱组件"],
    ])

    doc.add_heading("三、后端源码：鉴权、编排、服务与数据访问", level=1)
    doc.add_paragraph("后端使用 Express。它不实现向量模型本身，而是统一前端请求、调用 Flask 和 Neo4j、组织提示词，并把结果转为前端契约。")
    add_file_catalog(doc, "3.1 后端入口、类型和配置", [
        ["backend/package.json", "依赖与脚本", "npm 命令", "dev/build/start", "开发/构建"],
        ["backend/tsconfig.json", "TypeScript 编译配置", "源文件", "dist 编译规则", "构建"],
        ["src/types/index.ts", "后端共享类型", "无运行输入", "Chunk、KGPath、SpatialData、请求响应与实体类型", "编译期"],
        ["src/config.ts", "环境变量验证和配置对象", ".env", "端口、JWT、Neo4j、LLM、TopK 配置", "启动时"],
        ["src/index.ts", "Express 启动入口", "配置、路由、中间件", "监听 :3000 的 API 服务", "后端进程"],
    ])
    add_module(doc, "3.2 index.ts：服务启动", "注册所有中间件与路由，并初始化 Neo4j。", "环境配置、HTTP 请求。", "依次启用 CORS、JSON 解析、authMiddleware，挂载 qa/kg/admin/auth/upload/spatial 路由；提供 health 与 health/detailed；bootstrap 调 initNeo4j 后监听配置端口。", "统一的 /api 服务。", "前端只需访问一个 Node 服务，复杂依赖由该服务协调。", "backend/src/index.ts；backend/src/config.ts")
    add_module(doc, "3.3 JWT 中间件：middleware/auth.ts", "签发并验证身份令牌，保护一般 API 和管理员 API。", "登录成功后的 JwtPayload；请求 Authorization 头。", "signToken 以 JWT_SECRET 签名。authMiddleware 跳过登录和健康检查，其他请求解析 Bearer token 并验证；adminMiddleware 在此基础上检查 role=admin。", "通过时 req.user；失败时 401 或 403。", "路由权限不依赖前端隐藏按钮，后端仍会校验。", "backend/src/middleware/auth.ts")
    add_module(doc, "3.4 认证路由：routes/auth.ts", "验证演示账号并提供登录响应。", "POST /auth/login 的 username、password。", "从 DEMO_USERS 找账号，以 bcrypt.compare 比较密码哈希，成功后调用 signToken。", "token、username、role 或 400/401。", "前端可以获得后续受保护请求所需的身份。", "backend/src/routes/auth.ts")
    add_module(doc, "3.5 核心问答路由：routes/qa.ts", "协调 RAG、知识图谱、空间数据、LLM、SSE 和问答日志。", "POST /qa/ask 的 question、retrievalMode、stream。", "按 retrievalMode 决定 RAG/KG 是否参与；Promise.all 并行调用 services/tfidf.search、services/kg.searchEntities 和 getSpatialResults。优先用 Neo4j 真坐标，缺失时调用 matchMockSpatial。用 topScore<0.04 与 KG 命中量判断低相关性。流式模式写 meta/chunk/done，普通模式直接 JSON；两种模式都写 qa_logs。", "answer、sources、kgContext、spatialData，或 SSE 事件。", "并行降低延迟；低相关性关口避免无依据回答。", "backend/src/routes/qa.ts")
    add_module(doc, "3.6 Flask 检索客户端：services/tfidf.ts", "将 Node 后端的文献检索请求转发给 Python。文件名沿用 tfidf，但实际检索算法由 Flask 的 RETRIEVAL_MODE 决定。", "query、topK。", "POST :5000/search，映射 Python 返回的 chunk 字段；异常或非 2xx 时返回空数组。toSources 将 chunk 变成前端 Source。", "带 score 的 Chunk 列表、来源卡片数据。", "后端不用重复加载 BGE 模型，也能降级到只有 KG 的结果。", "backend/src/services/tfidf.ts")
    add_module(doc, "3.7 Neo4j 服务：services/kg.ts", "封装地质实体、子图、空间和管理 CRUD 查询。", "问题文本、关键词、跳数、Neo4j 节点 ID、更新字段。", "initNeo4j 建连接池。searchEntities 分词、过滤短词，对前 5 个关键词查关系并去重；getSubgraph 多跳展开；getSpatialResults 读取 lng/lat 和 path；管理函数读写实体并生成统计。每次会话最终关闭。", "KGPath、Subgraph、SpatialData、实体详情或统计。", "同一服务复用给问答、图谱、地图和管理台。", "backend/src/services/kg.ts；backend/src/services/tokenizer.ts")
    add_module(doc, "3.8 LLM 提示词与网关", "将本地检索证据组织成模型输入，并提供多模型与降级策略。", "问题、RAG chunks、KG paths、检索模式。", "llm.ts 的 buildPrompt 根据模式加入文档名/页码/相关度和关系链；gateway 按注册顺序调用 Provider。非流式的可重试错误会切换备用 Provider；无可用模型或失败时 mockAnswer 整理原文。流式模式使用首个 Provider，失败时一次性输出兜底答案。", "回答与来源，或 onChunk/onDone 回调。", "模型负责表达，文献和图谱负责提供上下文。", "backend/src/services/llm.ts；backend/src/services/llm/gateway.ts；backend/src/services/llm/provider.ts；backend/src/services/llm/deepseek.ts；backend/src/services/llm/tongyi.ts")
    add_module(doc, "3.9 SQLite 数据层：db/sqlite.ts", "读取 Python 管线产生的文档/切片，并维护问答日志。", "文档 ID、日志字段、统计请求。", "按需打开 geo_knowledge.db，设置 WAL，创建兼容的 qa_logs 表。getDocs/getDocById/getChunksByDocId 读知识库；insertQaLog 写本轮路径、命中数、耗时；getQaLogs/getStats 提供管理台数据。", "文档、切片、日志和计数。", "知识内容与问答运行记录存于同一 SQLite 文件但使用不同表。", "backend/src/db/sqlite.ts")
    add_module(doc, "3.10 管理、上传、图谱和空间路由", "向前端暴露各业务服务的 HTTP 接口。", "管理员 JWT、查询参数、JSON 请求体或 PDF FormData。", "admin.ts 保护 /admin 并转发文档/实体/日志操作；upload.ts 用 multer 校验 PDF 和大小后保存文件，后台 spawn Python 增量管线；kg.ts 提供 /kg/subgraph；spatial.ts 为 GeoJSON 图层预留接口，当前真实图层未接入。", "管理 JSON、上传结果、子图或空间接口状态。", "功能边界清晰：问答空间标注与独立 GeoJSON 图层接口可分别演进。", "backend/src/routes/admin.ts；backend/src/routes/upload.ts；backend/src/routes/kg.ts；backend/src/routes/spatial.ts")
    add_file_catalog(doc, "3.11 后端其余文件逐项说明", [
        ["services/llm/provider.ts", "Provider 接口", "统一参数和回调", "LLMProvider 契约", "DeepSeek/Tongyi/Gateway"],
        ["services/llm/deepseek.ts", "DeepSeek Provider", "系统提示词、用户提示词、chunks", "完整或流式回答、来源", "主模型"],
        ["services/llm/tongyi.ts", "通义千问 Provider", "同 Provider 接口参数", "完整或流式回答、来源", "可选备用模型"],
        ["services/llm/gateway.ts", "多模型调度与 mock", "Provider 列表、问答上下文", "模型结果或原文兜底结果", "问答服务"],
        ["services/tokenizer.ts", "Node 中文分词", "文本或文本数组", "过滤停用词后的 token", "KG 搜索"],
        ["data/spatial-mock.ts", "演示空间数据降级", "问题关键词", "矿点、岩石、构造线 SpatialData", "无真坐标时的 qa 路由"],
        ["test-qa.mjs", "问答接口调试", "本地服务与测试问题", "终端响应结果", "开发验证"],
    ])

    doc.add_heading("四、数据端源码：检索服务、建库和模型训练", level=1)
    doc.add_paragraph("数据端既包含常驻的 Flask 服务，也包含只在数据更新时执行的脚本。BERT-NER 与 OWL 推理已经接入离线链路；在线问答使用其写入 Neo4j 的结果，而不是对每个问题重新训练或重新推理。")
    add_module(doc, "4.1 配置与依赖", "让数据端各脚本共享数据库、PDF、模型、Neo4j 和 LLM 配置。", "config.py 中的本地真实配置；config.example.py 模板；requirements.txt。", "脚本从 config 导入 PDF_FILES、数据库路径、切片参数、RETRIEVAL_MODE、LLM_CONFIG、NEO4J_CONFIG。真实 config.py 被忽略，不应提交 API Key 或密码。", "每个脚本的运行参数和 Python 依赖环境。", "同一份配置可控制全管线和 Flask 行为。", "ml-service/config.example.py；ml-service/requirements.txt；ml-service/config.py（敏感本地文件）")
    add_module(doc, "4.2 Flask 服务：server.py", "提供 BGE/TF-IDF 文献检索、BERT 实体识别、本体查询、统计和开发面板。", "HTTP 请求、SQLite、BGE 模型、TF-IDF 模型、BERT-NER、Neo4j、OWL 文件。", "启动时加载 TF-IDF，按配置加载 BGE 索引，并加载 BERT-NER。/search 选择 _search_bge 或 _search_tfidf；/ner 运行 token 分类；/kg/search 查询图谱；/ontology/* 懒加载 OWL 后执行状态、推理或实体查询；/stats 汇总 SQLite 与 Neo4j；/start 与 /stop 管理开发服务进程。", "HTTP JSON，主要是检索 chunks 或实体/关系。", "后端在线必经的是 /search；/ner 和 /ontology 是已经可用的辅助接口，不是当前 /api/qa/ask 的必经调用。", "ml-service/server.py；ml-service/dashboard.html")
    add_file_catalog(doc, "4.3 Flask 端点逐项说明", [
        ["POST /search", "语义/关键词检索", "question, top_k", "chunks, scores", "Node services/tfidf.ts 在线调用"],
        ["POST /kg/search", "Python 侧子图查询", "keyword, limit", "nodes/relations", "开发或独立调用"],
        ["POST /ner", "BERT-NER 推理", "text", "entities", "辅助接口，未接入 qa.ts"],
        ["GET /health", "存活检查", "无", "status", "Node detailed health"],
        ["GET /stats", "数据库统计", "无", "Neo4j/SQLite 计数", "开发面板"],
        ["POST /start/<name>", "启动 Neo4j/Node/前端", "服务名", "进程状态", "开发面板"],
        ["POST /stop/<name>", "停止受管服务", "服务名", "进程状态", "开发面板"],
        ["GET /processes", "列出受管进程", "无", "运行状态/PID", "开发面板"],
        ["GET /dashboard", "返回开发面板页面", "无", "dashboard.html", "浏览器"],
        ["GET /ontology/status", "查看 OWL 概况", "无", "类/实例/属性数", "辅助验证"],
        ["POST /ontology/reason", "临时执行 HermiT", "无", "推理关系", "辅助验证"],
        ["POST /ontology/query", "按实体查 OWL", "entity", "类别与关系", "辅助验证"],
    ])
    add_module(doc, "4.4 总控离线管线：run_pipeline.py", "将 PDF 建库的步骤固定为可重复执行的顺序。", "config.PDF_FILES；--file 单 PDF；--from-step4。", "正常模式执行解析、切片、SQLite/TF-IDF、BGE、混合实体关系抽取、Neo4j 写入、OWL 推理。--file 临时替换 PDF_FILES 并启用增量；--from-step4 从现有 chunks 重组文本，跳过前段。", "终端统计、SQLite、图谱、本体等持久化结果。", "上传接口可把新 PDF 交给这一脚本后台处理。", "ml-service/scripts/run_pipeline.py")
    add_module(doc, "4.5 文档解析与切片", "将 PDF 转为可定位、可检索的小段。", "PDF 文件路径；解析结构；chunk_size/overlap。", "parse_pdf 用 pdfplumber 逐页提取并清洗文本。chunk_text 在每页文本中寻找优先断点，生成带 char_start/char_end 的重叠 chunk，避免生硬截断。", "解析文档对象和 chunk 列表。", "后端来源卡片能显示正确文档名和页码。", "ml-service/scripts/parse_pdf.py；ml-service/scripts/chunk_text.py")
    add_module(doc, "4.6 索引与语义向量", "让文档既可按关键词找，也可按语义相近找。", "所有 parsed_doc/chunks，或已有 SQLite。", "build_index 对中文分词，统一训练 TfidfVectorizer，建立 documents/chunks 表与 TF-IDF 模型。embed_chunks 读取 chunks 文本，SentenceTransformer 生成归一化 BGE 向量，存为 float32 BLOB，并记录 meta.vector_type/vector_dim。", "geo_knowledge.db、tfidf_model.pkl、512 维 chunk 向量。", "Flask 可由 RETRIEVAL_MODE 在 BGE 与 TF-IDF 间切换。", "ml-service/scripts/build_index.py；ml-service/scripts/embed_chunks.py")
    add_module(doc, "4.7 混合实体与关系抽取：extract_hybrid.py", "把地质报告由自然语言转为实体和关系。", "解析后文档、硬编码/Neo4j 词典、BERT-NER 模型、LLM 配置。", "第一步用预编译正则按 Mineral/Rock/Structure/TimePeriod/DepositType 扫描每页；第二步仅对有地质信号但词典覆盖不足的页面调用 BERT-NER，补入不在词典中的实体；第三步对含多个实体的页面并发调用 LLM 请求关系；RELATION_WHITELIST 检查实体类型对与合法关系，最后去重。", "concepts 和 relations。", "BERT-NER 不是只训练不使用：它是离线实体补漏的一环，产物随后进入 Neo4j。", "ml-service/scripts/extract_hybrid.py")
    add_module(doc, "4.8 Neo4j 图谱、本体与推理", "保存显式关系，补充可逻辑推导的隐式关系。", "实体、关系、文档结构；Neo4j；OWL 文件。", "write_neo4j 按地质标签 MERGE 节点，写 HOSTED_IN、CONTROLLED_BY、FORMED_IN 等边，并连接 Document-REFERENCES-Entity。build_ontology 从 Neo4j 读实体/关系构造 OWL 类、实例和对象属性。run_reasoning 保存推理前关系快照，运行传递关系/BFS 与本体推理，求差得到新增关系，写回 Neo4j 并标记 inferred=true。", "Neo4j 图、geo_planning.owl、推理关系。", "在线 KG 查询会读到 inferred 标记，因此 OWL 影响在线答案和图谱，但不会在每次问答实时启动推理器。", "ml-service/scripts/write_neo4j.py；ml-service/scripts/build_ontology.py；ml-service/scripts/run_reasoning.py")
    add_file_catalog(doc, "4.9 主建库和检索脚本逐项说明", [
        ["scripts/parse_pdf.py", "PDF 文本提取", "PDF path", "title/pages", "离线 Step 1"],
        ["scripts/chunk_text.py", "语义切片", "parsed doc,size,overlap", "chunks", "离线 Step 2"],
        ["scripts/build_index.py", "TF-IDF 与 SQLite", "文档/切片", "DB 与 pkl", "离线 Step 3"],
        ["scripts/embed_chunks.py", "BGE 编码", "SQLite chunks", "512 维向量", "离线 Step 3b"],
        ["scripts/extract_hybrid.py", "词典+BERT+LLM 抽取", "地质文本/模型", "实体关系", "离线 Step 4"],
        ["scripts/extract_concepts.py", "纯 LLM 概念关系抽取", "文本/LLM 配置", "concepts/relations", "规划类文档或备用"],
        ["scripts/extract_landuse.py", "用地分类提取", "解析文本", "层级节点", "特定文档"],
        ["scripts/write_neo4j.py", "图谱写入", "文档/实体/关系", "Neo4j", "离线 Step 5"],
        ["scripts/build_ontology.py", "OWL 构建", "Neo4j", "OWL/报告", "离线 Step 6"],
        ["scripts/run_reasoning.py", "本体推理并回写", "OWL", "inferred 关系", "离线 Step 6"],
        ["scripts/run_pipeline.py", "总控编排", "配置/CLI 参数", "全链路产物", "离线入口"],
    ])
    doc.add_heading("4.10 NER 训练链路：从标注数据到 bert-ner", level=2)
    doc.add_paragraph("这组脚本不参与每次在线问答。它们的任务是准备训练数据、微调模型、评估模型，然后由 extract_hybrid.py 与 Flask /ner 使用训练后的模型。")
    add_table(doc, ["阶段", "文件", "输入", "输出/作用"], [
        ["词典标注", "dict_ner_geology.py", "Neo4j 实体、内置术语、PDF", "正则生成 BIO/CoNLL，形成 ner_final.conll"],
        ["LLM 标注备选", "batch_ner_geology.py；batch_ner_annotate.py；batch_ner_smart.py；ner_annotate.py", "PDF 段落、LLM 配置", "多种 LLM 标注方案和 BIO 转换"],
        ["术语补充", "extract_glossary.py；extract_survey_terms.py；glossary_to_bio.py；merge_all_ner.py", "术语标准文本/CoNLL", "术语数据、BIO 语料和合并结果"],
        ["数据集准备", "prepare_ner_data.py", "ner_final.conll", "train/val/test CoNLL 切分"],
        ["模型训练", "train_ner.py", "训练/验证集", "models/bert-ner 模型和 tokenizer"],
        ["模型评估", "eval_ner.py", "模型、测试集", "类别指标和主力类 F1"],
    ], [2.3, 5.7, 4.5, 5.0])
    add_file_catalog(doc, "4.11 NER 与术语辅助脚本逐项说明", [
        ["dict_ner_geology.py", "词典正则标注主方案", "Neo4j/硬编码术语/PDF", "ner_final.conll", "训练数据准备"],
        ["batch_ner_geology.py", "LLM JSON 转 BIO", "PDF 缓存/LLM", "CoNLL", "备选标注"],
        ["batch_ner_annotate.py", "LLM BIO 直出批处理", "文本段落/LLM", "BIO 响应", "早期标注方案"],
        ["batch_ner_smart.py", "种子词筛选后标注", "分词段落/种子词", "较少 API 调用的标注", "备选标注"],
        ["ner_annotate.py", "LLM NER 通用工具", "文本/配置", "BIO/CoNLL", "被批处理脚本复用"],
        ["extract_glossary.py", "术语标准解析", "解析文本", "术语定义", "术语资料处理"],
        ["extract_survey_terms.py", "测绘术语解析", "标准 PDF 文本", "术语及同义词", "术语资料处理"],
        ["glossary_to_bio.py", "术语转 BIO", "术语列表", "CoNLL 文本", "NER 数据补充"],
        ["merge_all_ner.py", "多来源 CoNLL 合并", "多个 CoNLL 文件", "去重后的 CoNLL", "NER 数据整理"],
        ["prepare_ner_data.py", "训练集划分和清洗", "CoNLL", "train/val/test", "训练前"],
        ["train_ner.py", "BERT token 分类微调", "train/val", "bert-ner 权重", "模型训练"],
        ["eval_ner.py", "NER 评估", "模型/test", "F1 等指标", "训练验收"],
    ])
    add_file_catalog(doc, "4.12 数据端测试、环境与文档逐项说明", [
        ["test_parse_pdf.py", "解析结构测试", "测试 PDF", "断言结果", "手工验证"],
        ["test_chunk.py", "切片大小/重叠测试", "样本文本", "断言结果", "手工验证"],
        ["test_build_index.py", "SQLite 与 TF-IDF 测试", "临时数据", "断言结果", "手工验证"],
        ["test_extract_concepts.py", "LLM 解析工具单测", "模拟响应", "断言结果", "手工验证"],
        ["test_extract_landuse.py", "用地分类提取测试", "样本文本", "断言结果", "手工验证"],
        ["test_extract_glossary.py", "术语提取测试", "样本文本", "断言结果", "手工验证"],
        ["test_glossary_to_bio.py", "术语 BIO 测试", "术语样例", "断言结果", "手工验证"],
        ["test_ner_annotate.py", "NER 提示词/BIO 转换测试", "模拟数据", "断言结果", "手工验证"],
        ["test_write_neo4j.py", "Neo4j 读写测试", "在线 Neo4j", "节点/边验证", "集成验证"],
        ["test_bridge_landuse_concept.py", "分类-概念桥接测试", "名称样例", "断言结果", "历史/辅助验证"],
        ["verify_env.py", "Python 环境检查", "已安装依赖/模型", "可用性报告", "部署前"],
        ["CODE_INDEX.md", "数据端代码索引", "无", "脚本导航说明", "开发参考"],
        ["dashboard.html", "Flask 开发面板", "Flask API", "服务控制/状态页面", "开发运维"],
        ["scripts/owlDBDes.md", "OWL/图数据库说明", "无", "设计说明", "开发参考"],
    ])

    doc.add_heading("五、根目录入口、配置与数据产物", level=1)
    add_file_catalog(doc, "5.1 根目录文件", [
        ["start_all.py", "快捷启动 Flask 并打开开发面板", "本地 Python/浏览器", "Flask :5000 与 dashboard", "开发入口"],
        ["AGENTS.md", "项目结构、约定和常用命令", "无", "开发协作规则", "文档"],
        ["CONTEXT.md", "领域术语解释", "无", "Chunk、KG Path 等定义", "文档"],
        ["package.json", "工作区占位", "npm", "无实际运行脚本", "根目录配置"],
    ])
    add_file_catalog(doc, "5.2 生成数据与模型产物", [
        ["ml-service/output/geo_knowledge.db", "SQLite 知识库", "管线写入", "documents/chunks/vector/qa_logs", "在线检索和管理"],
        ["ml-service/output/tfidf_model.pkl", "TF-IDF 模型", "build_index 写入", "可加载词表/向量器", "Flask TF-IDF 模式"],
        ["ml-service/output/geo_planning.owl", "OWL 本体", "build_ontology 写入", "本体类/实例/关系", "推理和查询"],
        ["ml-service/models/bert-ner/", "训练后的 BERT-NER", "train_ner 写入", "config/model/tokenizer", "extract_hybrid 与 Flask /ner"],
        ["ml-service/output/*.conll", "NER 语料及切分", "标注/prepare 脚本", "训练/验证/测试语料", "模型训练"],
        ["backend/uploads/", "管理员上传 PDF", "upload 路由", "待/已处理 PDF", "增量建库输入"],
    ])

    doc.add_heading("六、按功能追踪：老师问到某项功能时该讲哪些文件", level=1)
    add_table(doc, ["老师可能问的问题", "应从哪里开始讲", "关键代码", "一句话回答"], [
        ["用户提问后为什么能边生成边显示？", "前端 SSE -> 后端 SSE", "api/qa.ts；stores/qa.ts；routes/qa.ts；llm/gateway.ts", "后端把模型 token 分段转成 SSE chunk，前端每收到一段就追加到同一条 AI 消息。"],
        ["回答依据来自哪里？", "RAG 文献检索", "server.py /search；build_index.py；embed_chunks.py；services/tfidf.ts", "PDF 已切成带页码的块并向量化，问题检索到最相近片段，再作为提示词上下文和来源卡片。"],
        ["知识图谱怎样参与回答？", "Neo4j KG", "extract_hybrid.py；write_neo4j.py；services/kg.ts；routes/qa.ts", "离线抽实体和关系写成图；在线按问题关键词查相关关系链并与文献一起交给模型。"],
        ["BERT-NER 是否真的使用？", "离线实体补漏", "train_ner.py；extract_hybrid.py；server.py /ner", "训练出的模型在混合抽取中对词典覆盖不足页面补充实体，也可通过 /ner 单独推理；当前问答不逐题调用 /ner。"],
        ["OWL 推理怎样影响系统？", "离线补全图谱", "build_ontology.py；run_reasoning.py；services/kg.ts", "推理得到的关系带 inferred 标记写回 Neo4j，在线 KG 查询能读到它们，但不会每个问题临时跑推理器。"],
        ["为什么要 Node 和 Python 两个后端？", "职责分层", "backend routes/services；ml-service/server.py", "Node 负责业务编排与鉴权；Python 负责向量模型、PDF、NER 和本体生态。"],
        ["管理员上传 PDF 后发生什么？", "异步增量建库", "AdminView.vue；upload.ts；run_pipeline.py", "文件先保存，后端后台启动 --file 管线，完成后新的 chunks 和图谱实体才可用于检索。"],
        ["地图数据是否全部是真实数据？", "空间数据边界", "services/kg.ts；data/spatial-mock.ts；routes/spatial.ts", "优先用 Neo4j 坐标；无坐标时问答页可用演示 mock，独立 GeoJSON 图层仍是预留接口，应如实说明。"],
    ], [4.0, 3.8, 6.0, 5.0])

    doc.add_heading("七、完整实例：一行代码到一轮答案", level=1)
    doc.add_paragraph("例题：用户以普通用户身份在问答页输入“尾亚钒钛磁铁矿受什么构造控制？”，选择混合检索。以下不是虚构具体地质结论，而是说明代码一定会执行的工作。")
    for text in [
        "LoginView/login store 已保存 JWT。router 守卫放行 /user/qa。",
        "ChatInput 发出文本和 hybrid 模式；QaView 调用 store.sendMessageStream。",
        "qa store 创建用户消息和空 AI 占位消息，调用 api/qa.askQuestionStream。",
        "浏览器带 Authorization 请求 POST /api/qa/ask，body 包含 question、stream:true、retrievalMode:hybrid。",
        "authMiddleware 验证 token；routes/qa.ts 并行调用 Flask 文献检索、Neo4j 关系检索、Neo4j 空间检索。",
        "Flask /search 使用 BGE 或 TF-IDF，从 SQLite 返回 Top-K 报告片段；kg.ts 以中文分词后的关键词查 Neo4j 路径。",
        "qa 路由先向前端发 meta。若两条证据均不相关，则发地质问题引导语；否则 gateway 构建“报告片段 + 图谱路径 + 问题”的提示词。",
        "DeepSeek Provider 逐段产生文字；qa 路由把每段封装成 SSE chunk。模型不可用时 gateway 把已检索原文组织为兜底回答。",
        "最终 onDone 带来源、图谱路径、空间数据；qa 路由写入 qa_logs。前端补齐来源卡片、关系小图和地图覆盖物。",
    ]:
        add_number(doc, text)

    doc.add_heading("八、当前实现边界", level=1)
    add_bullet(doc, "在线问答实际调用 Flask /search 和 Node 端 Neo4j 查询。BERT /ner 与 Flask /ontology/* 可独立使用，但没有被 routes/qa.ts 每次直接调用。")
    add_bullet(doc, "BERT-NER 和 OWL 都已接入离线建库：BERT 补实体，OWL 补关系；它们通过更新后的 Neo4j 间接提升在线问答。")
    add_bullet(doc, "独立 GeoJSON 空间图层接口仍是骨架；问答地图优先使用 Neo4j 属性，不足时可能使用演示 mock。")
    add_bullet(doc, "前端会话是 localStorage 持久化；后端 SQLite 保存的是问答统计日志，/api/qa/history 当前返回空数组。")
    add_bullet(doc, "流式 LLM 当前不做跨 Provider 续流，首个 Provider 失败时使用原文兜底；非流式模式才会切备用模型。")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()

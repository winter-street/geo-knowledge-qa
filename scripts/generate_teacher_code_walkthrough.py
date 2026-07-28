"""Generate the teacher-facing, code-by-code core-module walkthrough."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "地质找矿智能问答系统-核心模块代码讲解版.docx"
INK = "1A1A2E"
GREEN = "2E7D5B"
BLUE = "1E3A5F"
PALE = "E8F2EC"
GRAY = "666666"


def font(run, size=10.5, color=INK, bold=False, family="Microsoft YaHei"):
    run.font.name = family
    run._element.rPr.rFonts.set(qn("w:eastAsia"), family)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    tc_pr.append(element)


def borders(table, color="D6D2C8"):
    tbl_pr = table._tbl.tblPr
    element = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), color)
        element.append(node)
    tbl_pr.append(element)


def set_text(cell, value, bold=False, color=INK, size=9.2):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(value)
    font(run, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    borders(t)
    for i, item in enumerate(headers):
        shade(t.rows[0].cells[i], BLUE)
        set_text(t.rows[0].cells[i], item, bold=True, color="FFFFFF")
    for row in rows:
        cells = t.add_row().cells
        for i, item in enumerate(row):
            if i == 0:
                shade(cells[i], PALE)
            set_text(cells[i], item, bold=i == 0, size=8.8)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()


def paragraph(doc, text, indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(5)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    font(p.add_run(text))
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    font(p.add_run(text), size=10)


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    font(p.add_run(text), size=15 if level == 1 else 12, color=BLUE if level == 1 else GREEN, bold=True)


def code(doc, text):
    t = doc.add_table(rows=1, cols=1)
    borders(t, "BFC8D0")
    cell = t.cell(0, 0)
    shade(cell, "F5F7F8")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    font(run, size=8.8, color=INK, family="Consolas")
    doc.add_paragraph()


def callout(doc, title, text, color=GREEN):
    t = doc.add_table(rows=1, cols=1)
    borders(t, color)
    cell = t.cell(0, 0)
    shade(cell, "F6FAF7")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    font(p.add_run(title + "  "), size=10, bold=True)
    font(p.add_run(text), size=10)
    doc.add_paragraph()


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run("地质找矿智能问答系统 | 核心模块代码讲解版 | 2026 年 7 月"), size=8, color=GRAY)


def build():
    doc = Document()
    configure(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    font(p.add_run("地质找矿智能问答系统"), size=22, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    font(p.add_run("核心模块代码讲解版：RAG、知识图谱与 OWL 推理"), size=12, color=GRAY)
    callout(doc, "使用方式", "这份文档不是用户操作流程，而是答辩时的“打开代码讲解稿”。老师问一个模块，就按“问题 → 文件 → 函数 → 数据变化 → 在线收益 → 边界”顺序讲。")

    heading(doc, "一、先把项目的重点说准", 1)
    paragraph(doc, "我们的核心工作不是把大模型接到网页上，而是先把地质资料加工为两套可检索证据：一套是可回到页码的文献切片（RAG），一套是矿产、岩石、构造和年代之间的图谱关系（KG）。OWL 位于图谱的离线增强环节，它把已有、可信的关系按规则推出新关系，再物化回 Neo4j。在线问答只查询这些准备好的证据，并交给大模型组织表达。")
    table(doc, ["老师追问", "一句话回答", "立刻打开的代码"], [
        ["你们 RAG 怎么实现？", "不是只提关键词。离线把 PDF 按页切成重叠片段并向量化；在线将问题编码成同一向量空间，与所有切片算相似度，取 Top-K 原文和页码。", "chunk_text.py；embed_chunks.py；server.py"],
        ["知识图谱有什么用？", "它把跨句、跨页、跨文档重复出现的地质实体汇总为同一节点，让问题可沿“矿产—构造—年代—岩石”关系取证。", "extract_hybrid.py；write_neo4j.py；services/kg.ts"],
        ["OWL 有什么作用？", "它不凭空造地质事实，而是对具有传递性或明确定义的关系做规则推导，将原本要多跳访问的关系物化为 Neo4j 中的直接边。", "build_ontology.py；run_reasoning.py"],
    ], [4.0, 10.6, 4.8])

    heading(doc, "二、RAG：不是“提关键词然后搜索”，而是两阶段的文献证据检索", 1)
    heading(doc, "2.1 第一步：把 PDF 变成可定位的文本切片", 2)
    paragraph(doc, "先打开 ml-service/scripts/chunk_text.py 的 chunk_text()。它接收 parse_pdf.py 逐页提取的文本，不是把整本报告塞给模型，而是在每一页内按默认 512 字符窗口切片，并保留 128 字符重叠。重叠的目的，是让恰好落在切分边界处的“矿产—构造”语义不被拆散。")
    code(doc, "PDF -> parse_pdf(): [{page_num, text}]\n    -> chunk_text(): [{page, chunk_index, text, char_start, char_end}]\n    -> 每个切片仍保留文档名和页码，可在最终答案中溯源")
    table(doc, ["输入字段", "代码如何处理", "输出字段", "答辩要点"], [
        ["page_num、text", "长度小于 50 的页面跳过；短页整体保留；长页在目标窗口附近优先找段落、句号、换行、逗号断点。", "page、text、char_start、char_end、chunk_index", "切片不是机械截断，尽量在自然语言边界结束。"],
        ["chunk_size=512、overlap=128", "下一片起点为 end-overlap。", "相邻片段共享上下文。", "避免一句地质描述前半在片段 A、后半在片段 B 后无法单独理解。"],
    ], [3.2, 6.2, 4.3, 5.7])

    heading(doc, "2.2 第二步：为切片建立两种检索表示", 2)
    paragraph(doc, "打开 build_index.py 和 embed_chunks.py。系统保留 TF-IDF 作为词项检索和降级路径，同时以 BAAI/bge-small-zh-v1.5 生成 512 维语义向量作为默认语义检索路径。两种向量都写入 SQLite，但 BGE 向量会覆盖 chunks.vector 列，TF-IDF 的词表模型另存为 tfidf_model.pkl。")
    table(doc, ["路径", "离线代码", "做了什么", "适用价值"], [
        ["TF-IDF", "build_index()", "jieba 分词；对全部旧/新切片统一 fit_transform；保存词表模型和 tokens。", "词项精确匹配、BGE 不可用时降级。"],
        ["BGE", "embed_chunks()", "SentenceTransformer 将每个切片编码为归一化的 512 维 float32 向量，写回 SQLite BLOB。", "能召回用词不同但语义接近的地质描述。"],
    ], [2.2, 4.0, 6.5, 6.7])
    callout(doc, "老师问“提取关键词然后呢？”", "回答：TF-IDF 路径确实先用 jieba 把问题和切片分词，再计算 TF-IDF 向量余弦相似度；但默认 BGE 路径不以关键词匹配为终点，而是将完整问题编码为向量，与所有切片向量做点积。因为向量已归一化，点积等于余弦相似度。")

    heading(doc, "2.3 第三步：在线检索到底执行了什么", 2)
    paragraph(doc, "打开 ml-service/server.py 的 /search、_search_bge() 和 _search_tfidf()。Node 后端把问题和 Top-K 交给 Flask；Flask 根据 retrieval_mode 选择 BGE 或 TF-IDF。BGE 索引在服务启动时从 SQLite 一次性加载到内存，不是每次请求重新读完整数据库和重新编码全部文档。")
    code(doc, "BGE 查询：\nquestion -> '为这个句子生成表示以用于检索相关文章：' + question\n         -> model.encode(..., normalize_embeddings=True) 得到 q_vec\n         -> scores = chunk_vectors @ q_vec\n         -> argsort(scores)[::-1][:top_k]\n         -> [{doc_title, page, text, score}]\n\nTF-IDF 查询：\nquestion -> jieba 分词 -> vectorizer.transform(query)\n         -> 与每个 chunk.tokens 的向量计算 cosine_similarity\n         -> 排序取 Top-K")
    table(doc, ["RAG 的输入", "中间数据", "输出", "如何进入最终回答"], [
        ["问题：‘尾亚钒钛磁铁矿受什么构造控制？’；top_k", "BGE 查询向量 q_vec；内存中的 chunk_vectors；每条切片元数据（文档名、页码、文本、ID）。", "按 score 排序的 Top-K 切片。", "Node 的 toSources() 生成来源卡；buildPrompt() 将原文、页码、相关度写入 LLM 上下文。"],
    ], [4.0, 6.5, 3.6, 5.3])
    callout(doc, "RAG 的边界", "它返回的是‘相似的文本证据’，不是自动判定地质因果关系。因果/控制关系需要原文明确表述，或由图谱中的合法关系与 OWL 规则支持。")

    heading(doc, "三、从文献到 Neo4j：实体和关系不是直接由大模型随意生成", 1)
    paragraph(doc, "打开 ml-service/scripts/extract_hybrid.py。这里的策略是‘词典正则优先、BERT 补漏、LLM 只抽关系且受白名单约束’，目的是在地质术语识别的稳定性和覆盖率之间取得平衡。")
    table(doc, ["阶段", "函数/代码", "输入", "输出与约束"], [
        ["实体初筛", "_scan_page(text, patterns)", "预编译地质词典正则、单页文本。", "五类实体：Mineral、Rock、Structure、TimePeriod、DepositType；附带命中前后约 50 字上下文。"],
        ["BERT 补漏", "_extract_entities_bert(text)", "有地质信号且词典命中 1 到 9 个的页面。", "只补长度至少 3、词典中没有、类型合法的实体，避免对每页无差别跑模型。"],
        ["关系抽取", "_extract_relations_for_page()", "含至少两个已发现实体的页面、实体名到类型映射。", "LLM 只能从已发现实体中选；再由 RELATION_WHITELIST 检查类型对，例如 Mineral→Structure 只能是 CONTROLLED_BY。"],
        ["图谱落库", "write_geology_entities()、write_all()", "实体、合法关系、文档来源。", "MERGE 同名地质实体，写入关系；Document-REFERENCES-Entity 边保留文献溯源。"],
    ], [2.4, 4.5, 5.2, 7.3])
    callout(doc, "跨文献连接从哪里来？", "write_geology_entities() 对同名实体使用 MERGE，而不是每篇文章创建一个独立节点。两篇文献都提到同一构造或矿产时，会汇聚到同一节点，各自的 Document-REFERENCES 边保留出处。因此图谱可以把原本分散在不同文献中的关系串到同一实体周围。")

    heading(doc, "四、OWL：先把图谱读进本体，再按可解释规则推出新关系", 1)
    heading(doc, "4.1 先回答‘OWL 模块有什么作用？’", 2)
    paragraph(doc, "OWL 不是另一个检索库，也不是让大模型猜更多关系。它是对 Neo4j 已有实体、已有关系执行可检查的逻辑推导。它的价值是把需要多跳才能得到的、由规则必然成立的关系提前算出并回写 Neo4j；以后在线查询可以直接命中这条推理边，同时用 inferred=true 区分它和原始文本抽取的显式边。")
    code(doc, "显式图谱：早二叠世 -[BELONGS_TO]-> 二叠纪\n          二叠纪   -[BELONGS_TO]-> 古生代\n\n传递性推理：早二叠世 -[BELONGS_TO]-> 古生代\n\n若某矿物 -[FORMED_IN]-> 早二叠世，\n地学规则还可补出：该矿物 -[FORMED_IN]-> 二叠纪、古生代")
    callout(doc, "注意措辞", "上述例子中的新边来自‘年代层级’和‘形成时代向父级年代传播’规则，不是系统凭空断言某矿物的成因。所有推理事实都能回溯到已有边和规则。")

    heading(doc, "4.2 打开 build_ontology.py：它如何把 Neo4j 变为 OWL", 2)
    table(doc, ["阅读顺序", "代码做的事", "为什么需要"], [
        ["1. _read_entities() / _read_relations()", "从 Neo4j 读取五类地质节点和已有关系。", "OWL 的输入就是已经审核/抽取出的图谱事实。"],
        ["2. GEO_LABELS / RELATION_TYPES", "动态创建 Mineral、Rock、Structure、TimePeriod、DepositType 类和对象属性。BELONGS_TO、LIES_IN 被声明为 TransitiveProperty。", "把地质图谱中的标签和边映射为可以推理的本体语义。"],
        ["3. equivalent_to 定义", "定义 RockHostedMineral、StructurallyControlledMineral、AgeConstrainedMineral：一个矿物只有同时满足实体类型和支持关系，才属于对应语义类。", "得到的是可解释分类，而非按名称硬猜。"],
        ["4. ERA_HIERARCHY", "写入早/中/晚世→纪→代→宙的年代层级 BELONGS_TO 边。", "为跨层级年代关系提供可推导的基础。"],
        ["5. 保存 geo_planning.owl", "将实例、显式关系、年代层级写为 OWL 文件。", "使推理输入可复现，也便于独立检查。"],
    ], [4.3, 8.4, 6.7])

    heading(doc, "4.3 打开 run_reasoning.py：推理和回写如何发生", 2)
    table(doc, ["函数/步骤", "实际代码行为", "产生的结果"], [
        ["_snapshot_relations(onto)", "记录推理前的全部对象属性关系，形成 explicit 集合。", "有了基线，后面只能把差集认定为新推理关系。"],
        ["_infer_transitive_closure()", "对 BELONGS_TO、LIES_IN 建邻接表，用 BFS 求可达闭包。", "将 A→B→C 变成 A→C；这是减少后续多跳访问的直接来源。"],
        ["_apply_geology_rules()", "补 ASSOCIATED_WITH 反向边；将 Mineral-FORMED_IN-子年代传播到其父级年代。循环至没有新边。", "将保守、明确的地学规则落实为关系。"],
        ["sync_reasoner()", "调用 HermiT 进行 OWL 类层级推理；异常时不会丢弃前面已完成的传递闭包。", "得到等价类定义下的分类结果。"],
        ["all_relations - explicit", "重新收集关系并与基线做差。", "只保留本轮新增的 inferred 关系，而不重复标记旧边。"],
        ["write_to_neo4j()", "对每条新边 MATCH 两端节点，再 MERGE 关系并写 inferred:true；同时支持写 owlTypes 分类。", "Neo4j 中有可查询、可显示、可追溯的推理关系。"],
    ], [4.1, 9.1, 6.2])
    code(doc, "run_reasoning():\n  explicit = _snapshot_relations(onto)\n  inferred = transitive_closure(BELONGS_TO, LIES_IN)\n  inferred += apply_geology_rules(explicit | inferred) until stable\n  sync_reasoner()\n  inferred += all_relations_after_reasoner - explicit\n  return {relations: inferred, classifications: ...}\n\nwrite_to_neo4j(inferred_rels, classifications):\n  MERGE (a)-[:REL {display: ..., inferred: true}]->(b)")

    heading(doc, "4.4 这对在线检索有什么实际收益", 2)
    table(doc, ["没有推理边时", "推理边物化后", "代码中的在线表现"], [
        ["要从‘某矿物形成于早二叠世’得知‘与古生代有关’，需要查询矿物→早二叠世→二叠纪→古生代的多跳路径，并在应用层判断。", "Neo4j 已有矿物→古生代的 inferred: true 边；一跳关系查询即可返回。", "backend/src/services/kg.ts 的 searchEntities() 返回 r.inferred；LLM 提示词会说明该关系由 OWL 规则推得。"],
        ["不同文献各自提到同一实体或年代层级时，关系分散、查询上下文更长。", "同名实体 MERGE 到共同节点，规则再补足可推导边，跨文献证据可在同一局部子图汇合。", "在线 KG 路径和图谱页面都能读到该图中的关系。"],
    ], [6.3, 6.3, 6.8])
    callout(doc, "性能表述必须严谨", "可以说‘物化推理边有机会把某些多跳语义查询简化为一跳查询，从而减少运行时遍历’，不能说‘已经证明性能提升多少’，因为当前仓库没有保存推理前后响应时间的基准测试。")

    heading(doc, "五、当前代码边界：老师追问时要如实说明", 1)
    table(doc, ["事实", "如何准确说明"], [
        ["OWL 关系回写", "run_pipeline.py 会执行 build_ontology()、run_reasoning()，随后调用 write_to_neo4j(reason_stats['relations'])，因此新推理关系会以 inferred:true 回写 Neo4j。"],
        ["OWL 分类标签回写", "write_to_neo4j() 已支持第二个 classifications 参数并可写 n.owlTypes；独立运行 run_reasoning.py 时会传入它。但当前总管线调用只传 relations，所以总管线不会自动把本轮分类写入 owlTypes。这是当前调用方式的边界，不能说成已在全管线自动生效。"],
        ["关系的可信范围", "传递闭包仅用于 BELONGS_TO、LIES_IN；地学规则只做 ASSOCIATED_WITH 反向和 FORMED_IN 向父级年代传播。系统不会根据距离、名称相似或模型猜测生成新成矿因果。"],
        ["RAG 与 KG 的关系", "RAG 召回原文片段，KG 召回实体关系。两者并行提供给 LLM，但不是由 RAG 自动把每个片段转成图谱，也不是由 KG 替代文献引用。"],
    ], [4.5, 14.9])

    heading(doc, "六、老师当场问时可以这样回答", 1)
    table(doc, ["提问", "建议回答（可直接复述）"], [
        ["‘你们 RAG 是怎么做的？’", "我们先把 PDF 按页解析，用 512 字符窗口和 128 字符重叠切成可定位片段。离线用 BGE 把每段变成 512 维向量并写入 SQLite。用户提问时，把完整问题编码为查询向量，与内存中的所有切片向量做点积，取相似度最高的 Top-K，再把原文、文档名和页码交给大模型。TF-IDF 是分词后的备用检索路径。"],
        ["‘提取关键词之后呢？’", "如果走 TF-IDF，关键词经 jieba 分词后转换成 TF-IDF 向量，再和每个切片的词向量算余弦相似度；如果走默认 BGE，完整问题直接编码为语义向量，所以即使用词不完全相同，也能召回语义接近的报告片段。"],
        ["‘OWL 有什么用？’", "我们先从 Neo4j 读取显式实体关系，构建 OWL 类、属性和年代层级；再对 BELONGS_TO、LIES_IN 做传递闭包，对形成时代做父级传播，并调用 HermiT 做类别推理。新增关系和原始关系做差后，以 inferred:true 写回 Neo4j。因此在线检索能直接用到原本需要多跳才能得到的关系。"],
        ["‘这会不会瞎推理？’", "不会按相似名称或模型猜测推理。代码只对明确声明的传递关系和两条保守地学规则推理；每条新边都有原始边和规则可追溯，并且标记 inferred:true。"],
    ], [4.7, 14.7])

    heading(doc, "七、答辩展示顺序", 1)
    table(doc, ["顺序", "打开文件", "展示什么", "一句解释"], [
        ["1", "chunk_text.py", "512/128、自然断点、页码字段", "先说明原始报告如何变成可引用证据。"],
        ["2", "embed_chunks.py 和 server.py:_search_bge", "编码、归一化、点积、Top-K", "说明 RAG 是语义检索，不只是关键词匹配。"],
        ["3", "extract_hybrid.py 和 write_neo4j.py", "词典优先、BERT 补漏、关系白名单、MERGE", "说明图谱事实从哪里来、如何控制质量。"],
        ["4", "build_ontology.py", "TransitiveProperty、等价类、年代层级", "说明推理规则如何被显式定义。"],
        ["5", "run_reasoning.py", "快照、BFS 闭包、规则、差集、inferred:true 回写", "说明不是口头说 OWL，而是有完整的可执行推理和回写。"],
        ["6", "backend/src/services/kg.ts 与 routes/qa.ts", "读取 inferred、RAG+KG 并行、提示词", "最后说明离线成果怎样真正用于在线回答。"],
    ], [1.2, 5.2, 7.4, 5.6])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()

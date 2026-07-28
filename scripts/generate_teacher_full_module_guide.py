"""Generate the complete teacher-facing module walkthrough from verified code paths."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "地质找矿智能问答系统-完整模块代码讲解版.docx"
INK = "1A1A2E"
GREEN = "2E7D5B"
BLUE = "1E3A5F"
PALE = "E8F2EC"
GRAY = "666666"


def set_font(run, size=10.5, color=INK, bold=False, family="Microsoft YaHei"):
    run.font.name = family
    run._element.rPr.rFonts.set(qn("w:eastAsia"), family)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    props.append(node)


def add_borders(table, color="D6D2C8"):
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), color)
        borders.append(node)
    props.append(borders)


def write_cell(cell, value, bold=False, color=INK, size=8.9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(value), size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    add_borders(table)
    for index, header in enumerate(headers):
        shade(table.rows[0].cells[index], BLUE)
        write_cell(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if index == 0:
                shade(cells[index], PALE)
            write_cell(cells[index], value, bold=index == 0)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    doc.add_paragraph()


def add_paragraph(doc, text, indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(5)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    set_font(p.add_run(text))


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run(text), size=15 if level == 1 else 12, color=BLUE if level == 1 else GREEN, bold=True)


def add_code(doc, text):
    table = doc.add_table(rows=1, cols=1)
    add_borders(table, "BFC8D0")
    cell = table.cell(0, 0)
    shade(cell, "F5F7F8")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(text), size=8.8, family="Consolas")
    doc.add_paragraph()


def add_callout(doc, title, text, color=GREEN):
    table = doc.add_table(rows=1, cols=1)
    add_borders(table, color)
    cell = table.cell(0, 0)
    shade(cell, "F6FAF7")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run(title + "  "), size=10, bold=True)
    set_font(p.add_run(text), size=10)
    doc.add_paragraph()


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("地质找矿智能问答系统 | 完整模块代码讲解版 | 2026 年 7 月"), size=8, color=GRAY)


def module(doc, number, title, purpose, files, input_text, steps, output_text, effect, boundary, answer):
    add_heading(doc, f"{number}. {title}", 1)
    add_table(doc, ["项目", "内容"], [
        ["这个模块解决什么", purpose],
        ["打开哪些文件", files],
        ["输入是什么", input_text],
    ], [3.4, 16.0])
    add_table(doc, ["代码阅读顺序", "函数/位置", "内部发生了什么", "中间数据变成什么"], steps, [2.4, 4.8, 7.2, 5.0])
    add_table(doc, ["输出", "实际效果", "边界或注意点"], [[output_text, effect, boundary]], [5.4, 6.7, 7.3])
    add_callout(doc, "老师问到这里时可以这样回答", answer, BLUE)


def build():
    doc = Document()
    configure(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    set_font(p.add_run("地质找矿智能问答系统"), size=22, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    set_font(p.add_run("完整模块代码讲解版：给首次接触项目的老师与答辩成员"), size=12, color=GRAY)
    add_callout(doc, "这份文档的定位", "它不是文件清单，也不是产品宣传。每一章都回答：模块到底做什么、数据怎么进入、代码按什么顺序执行、产生什么可验证结果、对后续模块有什么作用、当前有什么边界。")

    add_heading(doc, "一、先建立整体认识：系统有两条完全不同的运行链路", 1)
    add_table(doc, ["链路", "什么时候运行", "从什么开始", "最后产物", "为什么要分开"], [
        ["离线建库链路", "新增 PDF 或主动重建资料时", "PDF 报告、术语词典、模型、Neo4j", "SQLite 切片与向量、Neo4j 图谱、OWL 推理边", "解析和推理耗时，不应该让每位提问用户重复等待。"],
        ["在线问答链路", "用户每次提问时", "问题、JWT、已建好的 SQLite/Neo4j", "流式答案、来源页码、图谱关系、可选空间结果、日志", "只读已准备好的证据，响应快且可追溯。"],
    ], [3.2, 3.2, 4.2, 4.6, 4.2])
    add_code(doc, "离线：PDF -> 解析/切片 -> BGE 与 TF-IDF 索引 -> 实体/关系 -> Neo4j\n      -> OWL 本体 -> 规则推理 -> inferred:true 关系回写 Neo4j\n\n在线：用户问题 -> JWT 校验 -> 并行 RAG / KG / 空间 -> 提示词组装 -> LLM\n      -> SSE 逐段答案 + 来源 + 关系 + 地图数据 -> 前端展示与日志")
    add_heading(doc, "二、模块总览：应重点讲的 12 个模块", 1)
    add_table(doc, ["编号", "模块", "它在系统中的位置", "上游依赖", "下游使用者"], [
        ["1", "身份与权限", "决定谁可以访问哪些页面和接口", "账号、JWT 密钥", "全部受保护 API、前端路由"],
        ["2", "流式会话前端", "把问题发送出去并实时显示答案", "JWT、SSE 协议", "问答页面、来源卡、地图"],
        ["3", "问答总编排", "选择检索路径、并行取证、控制低相关性", "RAG/KG/空间服务", "LLM、日志、前端 SSE"],
        ["4", "RAG 文献检索", "从报告中找可引用的原文片段", "PDF 切片、BGE/TF-IDF 索引", "LLM 提示词、来源卡"],
        ["5", "实体与关系抽取", "把报告文本转为地质实体和合法关系", "文本、词典、BERT、LLM", "Neo4j 图谱、OWL"],
        ["6", "Neo4j 图谱查询", "从结构化关系中找地质关联", "实体/关系图", "问答、图谱页、空间分析"],
        ["7", "OWL 本体与推理", "把已有关系按规则补成新关系", "Neo4j 显式图", "Neo4j inferred 边、在线 KG"],
        ["8", "LLM 证据生成", "把证据组织成专业中文答案", "RAG/KG/空间上下文", "SSE 文字、来源"],
        ["9", "空间条件分析", "解析范围与属性条件，筛选/排序地图要素", "坐标点线、查询条件", "地图、时空摘要、演示评分"],
        ["10", "图谱可视化", "把局部图谱变为可点击关系网络", "Neo4j 子图", "G6 图与节点详情"],
        ["11", "管理与资料上传", "维护文档、实体、日志和增量建库入口", "管理员 JWT、PDF", "后台管线、管理台统计"],
        ["12", "SQLite 资料与运行日志", "存文献切片、配置和问答记录", "建库结果、问答统计", "RAG、管理台、运行追踪"],
    ], [1.2, 3.2, 4.5, 4.8, 4.5])

    module(doc, "三", "身份与权限：前端不只是隐藏按钮，后端也会验证角色",
        "将普通用户与管理员分开，并避免仅靠前端页面隐藏来保护管理接口。",
        "frontend/src/stores/user.ts；frontend/src/router/index.ts；frontend/src/api/auth.ts；backend/src/routes/auth.ts；backend/src/middleware/auth.ts。",
        "登录页的 username/password；浏览器 localStorage 中的 token、username、role；请求头 Authorization。",
        [
            ["1", "routes/auth.ts：POST /auth/login", "从 DEMO_USERS 找演示账号，bcrypt.compare 比对密码哈希。", "账号密码 -> {token, username, role} 或 400/401。"],
            ["2", "middleware/auth.ts：signToken / authMiddleware", "用 JWT_SECRET 签发 24 小时 token；除 login/health 外解析 Bearer token，并把 payload 放到 req.user。", "HTTP 头 -> 已认证用户，或 401。"],
            ["3", "adminMiddleware", "读取 req.user.role，只允许 admin 访问 /api/admin 及上传路由。", "用户身份 -> 允许继续或 403。"],
            ["4", "stores/user.ts 与 router guard", "前端持久化 token/role；访问 /user、/admin 时同时检查是否登录和角色是否匹配。", "浏览器状态 -> 正确页面或重定向。"],
        ],
        "登录响应和响应式用户状态；受保护接口中的 req.user。",
        "管理员不能仅通过改 URL 获得权限；后端仍会作角色校验。",
        "当前账号是答辩演示账号，不是数据库用户体系；开发环境在完全无法连接后端时有本地登录回退，只能用于开发演示。",
        "我们有两层控制：前端路由控制界面入口，后端 JWT 和 adminMiddleware 控制真实 API 权限，因此不是把管理员按钮藏起来就算权限管理。")

    module(doc, "四", "流式会话前端：为什么答案能逐段出现且切换会话不丢失",
        "管理对话、检索模式和流式事件，将后端的多段 SSE 数据安全地更新为一条 AI 消息。",
        "frontend/src/components/ChatInput.vue；frontend/src/api/qa.ts；frontend/src/stores/qa.ts；frontend/src/views/QaView.vue；frontend/src/components/ChatMessage.vue。",
        "用户文本；当前 retrievalMode（rag/kg/hybrid）；JWT；后端 SSE 的 meta、chunk、done、error。",
        [
            ["1", "ChatInput -> QaView", "输入框 emit send，QaView 直接调用 store.sendMessageStream，并把当前检索模式双向绑定。", "文本/模式 -> store 方法。"],
            ["2", "sendMessageStream()", "先写入用户消息和内容为空的 AI 占位消息；若已有请求则调用 cancelStream 取消。", "会话列表立即出现本轮对话。"],
            ["3", "askQuestionStream()", "fetch POST /api/qa/ask，手动读取 response.body；按空行拆 SSE，再解析 data JSON。", "SSE 字节流 -> meta/chunk/done 回调。"],
            ["4", "onMeta/onChunk/onDone", "meta 可先挂空间数据；chunk 累加文本；done 补齐来源、KG 路径和空间分析，并持久化。", "空 AI 消息 -> 完整可追溯消息。"],
            ["5", "ChatMessage/QaView", "Markdown 渲染回答，展示 SourceCard、关系标签与地图入口。", "结构化结果 -> 用户可阅读的页面。"],
        ],
        "本地持久化的会话、逐段 AI 内容、来源/关系/空间字段。",
        "用户不用等待整段生成；消息中的来源和关系不会在流式过程中丢失。",
        "会话内容存于浏览器 localStorage；后端的 qa_logs 记录的是统计日志，不是完整聊天历史。",
        "SSE 不像普通 JSON 要等服务端生成完才返回。我们先插入占位消息，随后每收到一个 chunk 就替换该消息内容；done 才把来源、图谱和地图数据作为权威结果写齐。")

    module(doc, "五", "问答总编排：一条问题怎样决定走哪些证据路径",
        "统一处理一次问答：校验问题、按模式启用 RAG/KG、并行取证、低相关性拦截、调用 LLM、记录日志、输出 SSE 或 JSON。",
        "backend/src/routes/qa.ts；backend/src/services/runtime-settings.ts；backend/src/types/index.ts。",
        "POST /api/qa/ask 的 question、stream、retrievalMode；当前运行配置中的 Top-K。",
        [
            ["1", "读取请求参数", "去除 question 空白；缺失则 400；retrievalMode 默认 hybrid。", "HTTP body -> 合法问题与模式。"],
            ["2", "shouldRag / shouldKg", "rag 只查文献，kg 只查图谱，hybrid 两路都启用。空间解析独立执行，以识别空间意图。", "模式 -> 要执行的服务集合。"],
            ["3", "Promise.all", "RAG search、searchEntities、resolveSpatialForQuestion 同时启动。", "一个问题 -> 三类证据，降低串行等待。"],
            ["4", "isLowRelevance", "根据 RAG 最高分、KG 命中数、空间要素数判断。都没有依据时直接返回地质问答引导语。", "低相关问题 -> 不调用 LLM。"],
            ["5", "generateAnswer / generateAnswerStream", "将已取证结果交给网关；流式模式先写 meta，再写多个 chunk，最后写 done。", "证据 -> JSON 或 SSE。"],
            ["6", "insertQaLog", "记录问题、使用路径、命中文档数、KG/空间数量与耗时。", "本轮运行 -> SQLite 日志。"],
        ],
        "QaAskResponse 或 meta -> chunk* -> done 的 SSE 序列。",
        "RAG、KG、空间是并行，不是一个失败就导致整轮失败；低相关性关口减少无依据回答。",
        "路径开关决定是否把 RAG/KG放入回答上下文，不等于重新训练模型；空间功能只有在解析到空间意图或接口请求时返回分析结果。",
        "这个路由是后端调度中心。它本身不懂 PDF 或地质推理，而是协调已经准备好的检索、图谱和空间证据，再保证前端得到统一协议。")

    module(doc, "六", "RAG 文献检索：从 PDF 到 Top-K 原文证据的完整实现",
        "从地质报告中找与问题最相关、并能定位页码的文本片段，为答案提供可核验原文。",
        "ml-service/scripts/parse_pdf.py；chunk_text.py；build_index.py；embed_chunks.py；ml-service/server.py；backend/src/services/tfidf.ts。",
        "离线输入为 PDF；在线输入为 question、top_k 和 retrieval_mode（bge 或 tfidf）。",
        [
            ["1", "parse_pdf.py", "pdfplumber 逐页抽取并清理文本，保留 page_num。", "PDF -> [{page_num, text}]。"],
            ["2", "chunk_text()", "默认 512 字符窗口、128 字符重叠；优先段落/句子边界；记录字符范围。", "页文本 -> 可定位 chunks。"],
            ["3", "build_index()", "jieba 分词；所有切片统一训练 TfidfVectorizer；写 documents/chunks/tokens 和 tfidf_model.pkl。", "chunks -> 词项检索索引。"],
            ["4", "embed_chunks()", "BGE-small-zh-v1.5 对每个切片编码，归一化为 512 维 float32，写 SQLite vector BLOB。", "chunks -> 语义向量索引。"],
            ["5", "Flask /search", "按 retrieval_mode 选 _search_bge 或 _search_tfidf；BGE 在启动时加载向量和元数据到内存。", "问题 -> Top-K {title,page,text,score}。"],
            ["6", "services/tfidf.ts", "Node 调 Flask，并把 Python 字段转换为后端统一 Chunk；toSources 截取来源摘要。", "检索 JSON -> LLM 上下文和来源卡。"],
        ],
        "按相关度排序的文档片段，包含标题、页码、原文和 score。",
        "回答能引用报告原文；BGE 能召回用词不同但语义接近的段落，TF-IDF 保留关键词精确匹配与备用能力。",
        "RAG 只证明‘文献中存在相关文本’，不自动证明因果关系；相关性阈值用于拦截闲聊，但不是地质正确性的数学证明。",
        "我们不是提取几个关键词就全文查找。先把报告切片并编码；BGE 把完整问题编码为查询向量，与所有切片向量点积取 Top-K。由于向量已归一化，点积就是余弦相似度。TF-IDF 则是分词后做词向量余弦相似度的备用路径。")
    add_code(doc, "BGE：question -> encode(query_prefix + question) = q_vec\n     scores = chunk_vectors @ q_vec\n     top_indices = argsort(scores)[::-1][:top_k]\n     -> [{doc_title, page, text, score}]")

    module(doc, "七", "实体与关系抽取：怎样避免把大模型的自由生成直接写进图谱",
        "从报告文本抽取五类地质实体及其有明确证据的关系，并通过多层限制降低错误图谱边。",
        "ml-service/scripts/extract_hybrid.py；dict_ner_geology.py；entity_quality.py；ml-service/scripts/write_neo4j.py。",
        "单页文本；硬编码/图谱词典；BERT-NER 本地模型；LLM 配置。",
        [
            ["1", "hardcoded_entities / build_regex_patterns", "载入地质词典，预编译正则。", "词典 -> 可复用匹配模式。"],
            ["2", "_scan_page()", "逐页扫描 Mineral、Rock、Structure、TimePeriod、DepositType，并取命中前后约 50 字作为描述。", "页面 -> 稳定的候选实体。"],
            ["3", "_extract_entities_bert()", "只对有地质信号、词典覆盖 1 到 9 个实体且长度足够的页面补漏。", "候选实体 -> 少量新实体。"],
            ["4", "_extract_relations_for_page()", "只将已发现实体列表给 LLM；要求原文有依据；解析后按实体类型对检查 RELATION_WHITELIST。", "LLM JSON -> 合法关系或被过滤。"],
            ["5", "write_geology_entities()", "对合法类型的实体 MERGE；对关系两端 MATCH 后 MERGE 边。", "实体/关系 -> Neo4j 节点/边。"],
            ["6", "write_all()", "Document 节点与实体通过 REFERENCES 连接，保留哪篇报告出现了该实体。", "图谱事实 -> 文献溯源。"],
        ],
        "五类实体、合法关系、实体描述、Document-REFERENCES 溯源边。",
        "图谱不是纯 LLM 幻觉产物：词典优先，BERT 只补漏，关系必须通过类型白名单；同名实体 MERGE 后可连接多篇报告。",
        "词典覆盖外的新术语仍可能漏掉；关系白名单保证关系类型正确，但不能替代人工地质审核。",
        "我们让大模型只做它相对擅长的关系判断，而且限制它只能在已经识别出的实体中选择。随后再用类型白名单验证，例如矿产到构造只能是 CONTROLLED_BY，不匹配的关系不会写入 Neo4j。")

    module(doc, "八", "Neo4j 图谱查询：如何把实体关系变成在线问答证据",
        "按问题关键词在图谱中返回地质关系路径、局部子图、实体详情和带坐标的点线。",
        "backend/src/services/tokenizer.ts；backend/src/services/kg.ts；backend/src/routes/kg.ts；frontend/src/api/kg.ts。",
        "问题文本或关键词、跳数、实体类型、Neo4j 已建节点和关系。",
        [
            ["1", "tokenize()", "nodejieba 加载自定义地质词；过滤停用词、标点；不可用时降级字符切分。", "问题 -> 地质关键词列表。"],
            ["2", "searchEntities()", "取最多 5 个长度足够的关键词，Cypher 匹配包含关键词的节点及其关系，去重并限制数量。", "关键词 -> KGPath {from, relation, to, inferred}。"],
            ["3", "getSubgraph()", "先找种子节点，再在限定 depth 内展开路径，排除 Document 元数据节点。", "关键词/跳数 -> nodes + edges。"],
            ["4", "getSpatialResults()", "读取有 lng/lat 的节点和有 path 的构造线；同时携带年代、区域、OWL 类型等属性。", "问题 -> markers + polylines。"],
            ["5", "管理 CRUD", "getEntityDetail/updateEntity/deleteEntity/getEntityStats 使用 elementId 和属性更新。", "管理动作 -> 图谱实体状态。"],
        ],
        "KGPath、Subgraph、实体统计、带空间属性的点线数据。",
        "在线答案可补充‘矿产受何构造控制、形成于哪个年代、赋存在哪类岩石’这类结构化关系；图谱页能看到完整局部网络。",
        "searchEntities() 的常规路径以一跳关系为主并做数量限制；复杂全图推断不在用户请求时临时执行，而交给离线 OWL 推理预先物化。",
        "RAG 告诉我们报告里有哪些相关段落，Neo4j 告诉我们这些实体在图中的关系。我们把关系路径一起放进提示词，因此回答不只复述一段文本，还能解释实体之间的关联。")

    module(doc, "九", "OWL 本体与推理：为什么能补出原图没有的直接关系",
        "对 Neo4j 的显式地质关系执行可解释的传递与领域规则推理，并将新增关系物化回 Neo4j。",
        "ml-service/scripts/build_ontology.py；ml-service/scripts/run_reasoning.py；ml-service/scripts/run_pipeline.py；backend/src/services/kg.ts。",
        "Neo4j 中的五类实体、显式关系、年代层级；OWL 文件 geo_planning.owl。",
        [
            ["1", "build_ontology()", "_read_entities/_read_relations 从 Neo4j 读图；动态创建实体类和对象属性。BELONGS_TO、LIES_IN 声明为 TransitiveProperty。", "Neo4j 图 -> OWL 类、实例、属性。"],
            ["2", "equivalent_to 定义", "定义：Mineral 且 HOSTED_IN Rock -> RockHostedMineral；Mineral 且 CONTROLLED_BY Structure -> StructurallyControlledMineral 等。", "已有事实 -> 可解释语义分类。"],
            ["3", "ERA_HIERARCHY", "写入早/中/晚世 -> 纪 -> 代 -> 宙的 BELONGS_TO 层级。", "年代节点 -> 可传递层级链。"],
            ["4", "_snapshot_relations()", "推理前记录全部关系为 explicit 集合。", "原图 -> 可比较基线。"],
            ["5", "_infer_transitive_closure()", "对 BELONGS_TO、LIES_IN 构建邻接表并 BFS 求可达闭包。", "A->B->C -> 新 A->C。"],
            ["6", "_apply_geology_rules()", "补 ASSOCIATED_WITH 反向关系；矿物 FORMED_IN 子年代可传播到父级年代，直到没有新增边。", "显式+传递边 -> 保守领域推理边。"],
            ["7", "sync_reasoner() 与差集", "调用 HermiT；重新收集关系后与 explicit 取差；报告分类。", "本体结果 -> 只保留新增关系。"],
            ["8", "write_to_neo4j()", "MATCH 两端节点，MERGE 关系并设 inferred:true；函数也支持写 owlTypes 分类。", "推理结果 -> Neo4j 可查询边。"],
        ],
        "ontology_inference_report.json；Neo4j 中带 inferred:true 的关系；可选 owlTypes 分类。",
        "原本需要矿物->子年代->纪->代多跳才能查到的层级语义，可变为直接关系；在线 KG 查询可以标识该边来自推理。",
        "只对 BELONGS_TO/LIES_IN 做传递闭包，且只执行两条保守领域规则；当前 run_pipeline.py 回写 relations，但未把 classifications 参数传入，因此总管线不会自动写 owlTypes 分类标签。没有推理前后性能基准，不能量化宣称性能提升。",
        "OWL 的作用不是制造新地质发现。它先读取 Neo4j 的显式事实，再按传递性和年代层级等明确规则推导。例如早二叠世属于二叠纪、二叠纪属于古生代，就可补出早二叠世属于古生代；新增边会标 inferred:true 回写 Neo4j，在线查询可直接读到。")
    add_code(doc, "显式：早二叠世 -[BELONGS_TO]-> 二叠纪 -[BELONGS_TO]-> 古生代\n推理：早二叠世 -[BELONGS_TO {inferred:true}]-> 古生代\n\n若 矿物 -[FORMED_IN]-> 早二叠世，规则可补：\n矿物 -[FORMED_IN {inferred:true}]-> 二叠纪 / 古生代")

    module(doc, "十", "LLM 网关与提示词：为什么模型不会脱离检索证据直接回答",
        "将 RAG 原文、KG 路径和空间分析变成有固定约束的提示词，调用主/备模型并提供失败降级。",
        "backend/src/services/llm.ts；backend/src/services/llm/gateway.ts；backend/src/services/llm/deepseek.ts；backend/src/services/llm/tongyi.ts。",
        "问题、RAG chunks、KGPath、检索模式、空间分析结果；环境变量中的模型密钥和参数。",
        [
            ["1", "provider 注册", "有 DeepSeek key 时注册主 Provider；有通义 key 时注册备用 Provider。", "配置 -> 有序模型列表。"],
            ["2", "buildPrompt()", "按模式选择 RAG/KG；RAG 附文档名、页码、相关度；KG 附实体—关系—实体；空间附条件、结果和警告。", "结构化证据 -> 用户提示词。"],
            ["3", "SYSTEM_PROMPT", "要求基于上下文、信息不足时给已有线索、闲聊引导、空间演示数据须明确标注。", "模型行为约束。"],
            ["4", "gateway.generateAnswer()", "非流式按优先级调用 Provider；可重试错误尝试备用模型。", "证据 -> 完整 answer/sources。"],
            ["5", "gateway.generateAnswerStream()", "流式优先主 Provider；失败时以 mockAnswer 将已检索证据作为单段兜底输出。", "证据 -> chunk 回调与 done。"],
        ],
        "中文回答、来源；流式时为 onChunk/onDone 回调。",
        "模型承担表达，不承担资料库；服务异常时仍尽可能返回已检索到的文献或图谱线索。",
        "提示词约束能降低无依据回答，但不能保证绝对零幻觉；流式模式不在中途切换备用模型续流，而是降级为证据整理文本。",
        "我们不把问题直接丢给大模型。先把具体页码原文、图谱关系和空间警告装进提示词。模型的工作是把这些证据组织成中文答案；模型服务失败时，网关仍能把已检索原文返回给用户。")

    module(doc, "十一", "空间条件分析：如何把‘附近、范围内、按年代/评分排序’变成可计算结果",
        "将自然语言或结构化空间条件解析为实体类型、矿种、年代、中心、半径和排序，再筛选并分析点线要素。",
        "backend/src/routes/spatial.ts；backend/src/services/spatial-query.ts；backend/src/services/spatial.ts；backend/src/data/spatial-mock.ts；frontend/src/views/MapView.vue。",
        "question/keyword；可选 center、radiusKm、bbox、entityTypes、sortBy；Neo4j 坐标点线与可选模拟数据。",
        [
            ["1", "parseSpatialQuestion()", "正则识别矿种、实体类型、地质年代、距离半径、锚点名称、排序意图。", "自然语言 -> SpatialQueryInterpretation。"],
            ["2", "analyzeSpatialQuery()", "并行取 Neo4j 坐标与 mock；合并去重；若指定中心/锚点，构建距离参考。", "候选要素 -> 可分析点线。"],
            ["3", "distanceKm / pointToPathKm", "Haversine 计算点间距离；投影后计算点到折线段最短距离。", "坐标 -> 距离 km。"],
            ["4", "过滤与汇总", "按类型、矿种、年代、成因、OWL 类型、bbox、半径过滤；按时间和区域汇总。", "候选集 -> 满足条件的要素和统计。"],
            ["5", "scoreMineralProspectivity()", "组合构造邻近度、OWL/KG 证据、地质属性完整度、同类矿点聚集度，按 score 排序。", "矿点 -> demo-v1 规则评分和因素。"],
            ["6", "routes/spatial.ts / MapView", "接口转为 GeoJSON feature；高德地图渲染；前端可按时代/区域过滤并导出 CSV/GeoJSON。", "分析 JSON -> 地图与文件。"],
        ],
        "markers/polylines、距离、区域/时代摘要、警告、规则评分、GeoJSON/CSV 导出。",
        "使空间化问题有透明的条件解析和计算过程，而不是只显示静态点位。",
        "成矿有利度是 demo-v1 规则模型，只用于功能演示和方案比较，不是经训练验证的预测模型；含 mock 的结果会标出不可作为实际勘查结论；独立 GeoJSON 地质图层接口仍待真实数据接入。",
        "空间模块的关键不是画地图，而是把‘尾亚附近 20 千米、钒钛磁铁矿、按评分排序’拆成可计算条件，再计算距离、筛选要素并给出每个评分因子的依据。它目前是规则型分析，不是找矿预测模型。")

    module(doc, "十二", "图谱可视化：如何把 Neo4j 的关系变成可理解的网络",
        "让用户按关键词和跳数查看局部知识图谱，并通过颜色、同心环和节点详情理解关系网络。",
        "frontend/src/views/GraphView.vue；frontend/src/api/kg.ts；backend/src/routes/kg.ts；backend/src/services/kg.ts。",
        "搜索关键词、hopCount、后端 Subgraph 的 nodes/edges。",
        [
            ["1", "loadGraph()", "调用 getSubgraph(keyword, hopCount)，将返回 nodes/edges 写入响应式状态。", "关键词 -> 局部图。"],
            ["2", "computeHopDistances()", "对主种子节点 BFS，计算每个节点到种子的最短跳数。", "边集合 -> hop distance。"],
            ["3", "renderGraph()", "按 hop 分组放到同心圆，随后使用 G6 力导向布局；实体类型决定颜色。", "图数据 -> 布局后的画布。"],
            ["4", "节点点击与 OWL 标识", "点击节点显示属性与关联边；带 owlTypes 的节点可显示推理分类提示。", "节点 -> 可解释详情。"],
        ],
        "可点击 G6 图、实体统计、节点详情。",
        "将问答中的一条关系扩展为可观察的局部网络，便于老师或用户理解知识图谱不是一张静态图片。",
        "展示的是关键词周围限定跳数的子图，不是全库全量图；跳数越大，节点数和渲染复杂度越高。",
        "图谱页不是另建一套数据，它直接消费 Neo4j 子图接口。前端先用 BFS 计算与中心实体的跳数，再按同心环和力导向布局，让一跳、二跳关系直观可见。")

    module(doc, "十三", "管理台与 PDF 上传：资料怎样进入可重复的建库流程",
        "让管理员查看知识库规模、文档切片、实体和问答日志，并以受控方式上传 PDF、触发增量处理。",
        "frontend/src/views/AdminView.vue；backend/src/routes/admin.ts；backend/src/routes/upload.ts；ml-service/scripts/run_pipeline.py。",
        "管理员 JWT；PDF FormData；管理员查看/编辑实体、查看日志的请求。",
        [
            ["1", "adminMiddleware", "所有 /admin 路由先检查 role=admin。", "请求 -> 管理权限。"],
            ["2", "admin.ts", "读取 SQLite 文档/切片/日志统计，调用 KG 服务读取、修改、删除实体。", "数据库/图谱 -> 管理 JSON。"],
            ["3", "upload.ts：multer", "仅允许 application/pdf，限制 50 MB；以时间戳+原始中文名存入 backend/uploads。", "上传流 -> 本地 PDF 文件。"],
            ["4", "spawn(run_pipeline.py --file)", "后台、脱离当前 HTTP 响应地启动单文件增量管线。", "PDF 路径 -> 离线建库任务。"],
            ["5", "run_pipeline.py", "单文件模式依次完成解析、切片、索引/向量、实体关系抽取、Neo4j 写入、OWL 推理。", "新资料 -> 可被在线检索的知识。"],
        ],
        "管理统计、文档列表、切片/实体详情；上传已接收响应；后台建库任务。",
        "管理员无需手动打开 SQLite 或 Neo4j 即可维护资料，新增资料能走同一套数据加工流程。",
        "上传接口立即返回‘后台任务已启动’，不等待管线完成，也没有在接口中回传最终成功/失败状态；实际运行结果仍需通过日志、管理统计或 Python 控制台确认。",
        "上传不是把 PDF 放进一个文件夹就完成。接口校验类型和大小后，把文件交给同一条 run_pipeline 增量管线，因此新资料会经历和初始资料相同的切片、向量、实体、图谱和推理步骤。")

    module(doc, "十四", "SQLite 资料库与日志：哪些数据放在 SQLite，哪些数据放在 Neo4j",
        "将大段文献及其检索信息放入 SQLite，将实体关系放入 Neo4j，并将运行统计记录下来。",
        "backend/src/db/sqlite.ts；ml-service/scripts/build_index.py；backend/src/services/runtime-settings.ts。",
        "离线产生的 documents/chunks；问答记录；管理台统计请求；运行配置。",
        [
            ["1", "documents 表", "保存标题、类型、页数等文档元信息。", "文档 -> 可管理目录。"],
            ["2", "chunks 表", "保存 doc_id、page、text、tokens、vector、字符位置。", "页文本 -> RAG 可查切片。"],
            ["3", "qa_logs 表", "insertQaLog 记录问题、路径、命中数、KG/空间数量、耗时。", "每轮问答 -> 可统计日志。"],
            ["4", "sqlite.ts 查询函数", "getDocs/getChunksByDocId/getStats/getQaLogs 向管理台提供读接口。", "SQLite -> 管理界面数据。"],
            ["5", "runtime settings", "保存/校验 RAG 模式、Top-K、权重、LLM 参数等运行设置。", "管理员配置 -> 后端检索/生成行为。"],
        ],
        "文档与切片数据、检索索引元数据、问答统计、运行配置。",
        "SQLite 擅长存大量文本切片和顺序查询；Neo4j 擅长多实体关系与图遍历，职责清楚。",
        "chunks.vector 最终由 BGE 512 维向量覆盖；TF-IDF 的主要可复用产物是独立的模型文件和 tokens。qa_logs 不是完整聊天记录。",
        "数据分库不是重复存储：SQLite 保存适合检索的文本证据和运行记录，Neo4j 保存适合关系查询的实体网络。问答时两者各做擅长的部分。")

    add_heading(doc, "十五、最终串联：一个新人应该怎样理解系统的因果关系", 1)
    add_table(doc, ["先有", "因此产生", "随后被谁使用", "最终对用户意味着什么"], [
        ["PDF 被解析、切片、向量化", "SQLite 中出现带页码的切片和 BGE/TF-IDF 索引", "Flask /search 与 RAG 客户端", "回答可给出报告原文和页码。"],
        ["文本被识别为实体并抽出合法关系", "Neo4j 中出现地质实体、关系和文献引用", "KG 查询、图谱页、空间模块、OWL", "回答可解释矿产/构造/岩石/年代关联。"],
        ["显式图被构造成 OWL 并运行规则", "新关系以 inferred:true 物化回 Neo4j", "在线 KG 查询和提示词", "可得到原图中需多跳才能发现的、规则可证明的关联。"],
        ["问答总编排并行拿到三类证据", "LLM 提示词带原文、关系、空间警告", "SSE 前端与来源/地图组件", "用户看到的是逐段答案，并能查看依据。"],
    ], [4.5, 5.3, 4.7, 5.0])
    add_callout(doc, "最重要的答辩结论", "系统的核心价值是把‘不可直接问答的 PDF 资料’变成可追溯文本证据、可查询关系证据和可解释推理证据；大模型只负责基于这些证据组织回答，而不是替代资料库或推理规则。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()

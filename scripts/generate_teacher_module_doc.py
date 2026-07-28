from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / "docs" / "地质找矿智能问答系统-模块说明（教师版）.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(9)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "1E3A5F")
        set_cell_text(cell, text, bold=True, color="FFFFFF")
        if widths:
            cell.width = Cm(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], str(text))
            if widths:
                cells[i].width = Cm(widths[i])
    doc.add_paragraph()
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_module(doc, title, purpose, inputs, process, outputs, effect=None, files=None):
    doc.add_heading(title, level=3)
    p = doc.add_paragraph()
    p.add_run("用途：").bold = True
    p.add_run(purpose)
    p = doc.add_paragraph()
    p.add_run("输入：").bold = True
    p.add_run(inputs)
    p = doc.add_paragraph()
    p.add_run("内部处理：").bold = True
    p.add_run(process)
    p = doc.add_paragraph()
    p.add_run("输出：").bold = True
    p.add_run(outputs)
    if effect:
        p = doc.add_paragraph()
        p.add_run("带来的效果：").bold = True
        p.add_run(effect)
    if files:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run("对应代码：")
        r.bold = True
        p.add_run(files)


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def setup_doc(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color in [("Title", 24, "1E3A5F"), ("Heading 1", 17, "1E3A5F"), ("Heading 2", 13, "2E7D5B"), ("Heading 3", 11.5, "1E3A5F")]:
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(13)
        style.paragraph_format.space_after = Pt(7)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("地质找矿智能问答系统｜模块说明（代码实现版）")
    set_run_font(r, size=8, color="6B6B78")


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(75)
    r = p.add_run("地质找矿智能问答系统")
    set_run_font(r, size=25, bold=True, color="1E3A5F")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("功能模块说明与运行过程（教师版）")
    set_run_font(r, size=17, bold=True, color="2E7D5B")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("依据当前项目源码整理｜用于答辩说明与教师查阅")
    set_run_font(r, size=10, color="6B6B78")
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(115)
    r = p.add_run("2026 年 7 月")
    set_run_font(r, size=10, color="6B6B78")
    doc.add_page_break()


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_doc(doc)
    add_title(doc)

    doc.add_heading("一、先用一句话说清系统", level=1)
    doc.add_paragraph(
        "这是一个面向地质找矿资料的问答系统。它把 PDF 地质报告加工成可检索的文本片段和实体关系图；"
        "用户提出问题后，系统同时查报告片段和知识图谱，再把查到的证据交给大模型组织成回答。"
        "回答除了文字，还可以带来源、实体关系和地图标注。"
    )
    doc.add_heading("1.1 为什么不是直接把问题交给大模型", level=2)
    doc.add_paragraph(
        "通用大模型不知道本项目导入了哪些报告，也无法保证回答能追溯到哪一页资料。"
        "因此系统先检索本地知识库，再让模型基于检索结果回答。这里的核心思想是："
        "模型负责把信息讲清楚，项目自己的文献和图谱负责提供依据。"
    )
    add_table(doc, ["组成", "运行位置", "主要职责", "对外产出"], [
        ["前端", "浏览器 / :5173", "让用户登录、提问、查看地图、查看图谱和管理资料", "页面、流式回答、图谱和地图展示"],
        ["后端", "Node.js Express / :3000", "鉴权、协调检索、调用模型、记录日志、提供管理接口", "统一的 API 和 SSE 事件流"],
        ["数据端", "Python Flask / :5000", "运行时文献检索；离线将 PDF 加工为知识库", "相关文本片段、索引、实体关系、本体文件"],
        ["存储层", "SQLite + Neo4j + 文件", "保存文档切片、向量、问答日志、实体和关系", "可被检索和可视化的数据"],
    ], [2.0, 3.0, 7.0, 5.0])

    doc.add_heading("1.2 两条时间线：先加工资料，再回答问题", level=2)
    add_table(doc, ["时间线", "什么时候运行", "做的事情", "得到什么"], [
        ["离线数据加工", "导入新 PDF 或重建知识库时", "解析 PDF、切片、建立检索索引、抽取实体关系、写入图谱和本体", "SQLite 文档库、Neo4j 图谱、OWL 本体"],
        ["在线问答", "用户每次发送问题时", "鉴权、并行检索文献和图谱、调用大模型、把回答流式推回页面", "回答文字、来源卡片、关系数据、空间标注和日志"],
    ], [3.0, 4.0, 6.0, 6.0])
    doc.add_paragraph("答辩时可以这样概括：离线链路相当于整理资料和建索引；在线链路相当于根据索引快速查资料并作答。", style="Intense Quote")

    doc.add_heading("二、前端：用户看到和操作的部分", level=1)
    doc.add_paragraph(
        "前端采用 Vue 3。它本身不做地质知识判断，主要负责收集用户操作、调用后端接口、保存界面状态，"
        "并把返回的数据以聊天、地图和关系图的形式呈现。普通用户和管理员共用问答、地图、图谱页面；"
        "管理员额外拥有知识库管理页面。"
    )
    add_table(doc, ["前端模块", "用户输入", "页面输出"], [
        ["登录与路由", "账号、密码；访问页面地址", "登录状态；按 user/admin 角色进入对应页面"],
        ["智能问答", "问题、检索模式、发送/取消操作", "逐段增长的回答、来源、关系小图和空间数据"],
        ["地图交互", "底图、绘制和图层操作；问答返回的坐标", "高德地图、点标注、线标注或文字降级结果"],
        ["知识图谱", "关键词、展开跳数、节点点击", "可拖拽缩放的节点-边关系图和节点详情"],
        ["管理控制台", "上传 PDF、查看/删除文档、编辑实体", "文档统计、切片信息、问答日志和实体管理结果"],
    ], [3.2, 6.1, 9.7])

    add_module(doc, "2.1 登录、角色与路由守卫", "控制谁能进入系统，以及不同角色能看到哪些页面。", "用户名、密码；浏览器保存的 token 和 role；用户准备访问的路由。", "登录页把账号密码发送给后端。登录成功后，token、用户名和角色写入浏览器 localStorage。路由守卫在每次跳转前检查 token：未登录转到登录页；角色不匹配则转回该角色首页。", "普通用户进入问答、地图、图谱；管理员进入控制台，并可访问管理接口。", "把页面权限和后端接口权限都做了基本隔离。", "frontend/src/views/LoginView.vue；frontend/src/router/index.ts；frontend/src/stores/user.ts")
    add_module(doc, "2.2 问答状态仓库：保存一轮对话", "把输入框、聊天消息列表和流式结果连接起来，是问答页面的前端调度中心。", "用户问题；当前检索模式（RAG、KG、混合）；后端推回的 meta、chunk、done 事件。", "发送后，先向消息列表加入用户消息和一个空的 AI 消息，界面无需等待。随后调用流式接口；每收到一段 chunk，就把该段文字追加到空消息中。done 事件到达时，再补上来源、图谱关系和空间数据。对话记录写入 localStorage。", "随时间更新的消息列表、加载状态和错误信息。", "实现“边生成边显示”，并在刷新浏览器后保留本地对话。", "frontend/src/stores/qa.ts")
    add_module(doc, "2.3 流式问答接口：读取 SSE", "把后端的连续事件流转换为前端可用的回调。", "问题文本、检索模式、token，以及处理 meta/chunk/done/error 的回调函数。", "使用 fetch 发送 stream:true。开发环境直接连接 :3000，避免开发代理缓冲响应。浏览器逐行读取 data: 开头的 SSE 数据，解析 JSON 后按事件类型分发；同时过滤不完整的空间点和线，避免地图收到异常坐标。", "取消函数，以及持续触发的回调数据。", "前端不必理解网络流细节，问答仓库只需关心“收到一段文字”或“本轮完成”。", "frontend/src/api/qa.ts")
    add_module(doc, "2.4 问答页面与回答组件", "把对话列表、消息区、输入框、来源卡片、关系小图和内嵌地图组合成一次完整问答体验。", "问答仓库的消息；用户点击发送、新建/切换会话、查看地图等操作。", "QaView 监听消息变化。AI 消息中有 spatialData 时，页面会尝试加载高德地图并渲染标记点和折线；地图密钥不可用时改以文字列表展示。ChatMessage 负责 Markdown 回答、来源卡片和知识图谱小图。", "用户可读的回答及其证据入口；有坐标时可见地图标注。", "回答不只是一段文字，用户可以追问“证据来自哪里”“相关实体有哪些”。", "frontend/src/views/QaView.vue；frontend/src/components/ChatMessage.vue；frontend/src/components/SourceCard.vue；frontend/src/components/KgMiniGraph.vue")
    add_module(doc, "2.5 知识图谱页面", "单独查看某个实体周边的一到三跳关系，便于理解矿产、岩石、构造和年代之间的联系。", "搜索关键词，例如“磁铁矿”；跳数 1、2 或 3。", "页面向后端请求 nodes 和 edges。前端先找到与关键词最匹配的中心节点，再用广度优先搜索计算每个节点与中心节点的跳数，把不同跳数的节点放在不同圆环上，最后交给 AntV G6 渲染并支持拖拽、缩放、节点点击。", "关系网络图、各实体/关系类型统计、节点详情。", "把数据库里难以阅读的关系记录转为直观网络。", "frontend/src/views/GraphView.vue；frontend/src/api/kg.ts")
    add_module(doc, "2.6 地图页面与管理控制台", "地图页面服务于空间查看；管理控制台服务于资料和实体维护。", "地图操作、图层操作；管理员上传的 PDF、文档/实体操作。", "MapView 负责高德地图实例、绘制工具和问答空间数据的显示。AdminView 调用统计、文档、日志和实体接口；上传时发送 multipart/form-data，后端接收文件后异步启动数据加工脚本。", "地图图层和标注；统计卡片、文档清单、切片详情、日志、实体编辑结果。", "将知识库维护从命令行操作转到页面操作。", "frontend/src/views/MapView.vue；frontend/src/views/AdminView.vue；frontend/src/api/doc.ts")

    doc.add_heading("三、后端：把一次提问编排成完整流程", level=1)
    doc.add_paragraph(
        "后端采用 Node.js + Express，是各服务的统一入口。前端不会直接访问 Neo4j 或数据库，"
        "而是通过后端访问。这样既能做登录鉴权，也能把多路检索的结果统一成前端需要的格式。"
    )
    add_table(doc, ["后端模块", "输入", "输出", "核心作用"], [
        ["入口与鉴权", "HTTP 请求、Authorization 头", "放行后的请求或 401/403", "统一解析 JSON、跨域和身份验证"],
        ["问答路由", "问题、检索模式、是否流式", "JSON 回答或 SSE 事件流", "并行协调文献、图谱和空间结果"],
        ["检索客户端", "问题、Top-K", "带分数的文本片段", "调用 Flask /search"],
        ["知识图谱服务", "问题/关键词/节点 ID", "关系路径、子图、空间要素", "查询 Neo4j"],
        ["大模型网关", "问题、文献片段、关系路径", "回答文本和来源", "构建提示词、调用主备模型、失败降级"],
        ["SQLite 数据层", "文档/日志查询或写入参数", "记录、统计和列表", "读取知识库文档并记录问答日志"],
    ], [3.0, 4.0, 4.0, 7.5])

    add_module(doc, "3.1 入口、CORS 与 JWT 鉴权", "提供统一 API 入口，保护除登录和健康检查以外的接口。", "所有 /api 请求；Authorization: Bearer <token>。", "Express 先启用跨域和 JSON 解析，再执行 authMiddleware。中间件跳过 /api/auth/login 和 /api/health；其他请求必须验证 JWT。管理接口在普通验证后还会检查 role 是否为 admin。", "请求对象中附带用户信息，或返回未登录/无权限错误。", "避免把 Neo4j、文档管理等能力直接暴露给未登录访问者。", "backend/src/index.ts；backend/src/middleware/auth.ts")
    add_module(doc, "3.2 登录接口", "验证演示账号密码并签发 JWT。", "POST /api/auth/login 的 username 和 password。", "后端从内置演示用户表找到账号，使用 bcrypt 对比密码哈希；验证通过后使用 JWT 密钥签发包含 userId、username、role 的 token。", "{ token, username, role }，失败时为 400 或 401。", "前端随后可以带 token 调用受保护接口。", "backend/src/routes/auth.ts")
    add_module(doc, "3.3 问答路由：在线链路的总调度", "一次性完成检索、相关性判断、生成、流式返回和日志记录。", "POST /api/qa/ask：question、retrievalMode（rag/kg/hybrid）、stream。", "根据模式决定是否调用文献检索和图谱检索；同时，无论检索模式如何都会尝试读取空间数据。三项任务用 Promise.all 并行执行。若文献最高分低于 0.04 且图谱没有命中（或当前模式对应路径没有命中），返回地质问题引导语，不调用大模型。否则交给大模型网关。流式时依次发送 meta、若干 chunk、done；结束时写入问答日志。", "普通模式返回 answer、sources、kgContext、spatialData；流式模式返回 SSE 事件。", "并行减少等待时间；低相关性关口避免把闲聊硬解释成地质结论。", "backend/src/routes/qa.ts")
    add_module(doc, "3.4 文献检索客户端", "将后端需要的检索请求转交给 Python 服务，并把返回结构统一成后端 Chunk。", "问题文本、Top-K。", "通过 HTTP POST 调用 http://127.0.0.1:5000/search，传入 question 和 top_k。返回的 id、text、doc_title、page、score 被映射为后端统一结构；Flask 不可用时安静返回空数组，让问答路由仍可继续走图谱或降级回答。", "若干 {chunk, score}；以及由它转换出的来源卡片 Source。", "Node 端不必加载 Python 的向量模型，职责边界清楚。", "backend/src/services/tfidf.ts")
    add_module(doc, "3.5 知识图谱与空间查询服务", "从 Neo4j 读取与问题有关的实体关系、用于图谱页面的子图，以及实体坐标。", "问题文本、关键词、跳数或实体 ID。", "先用中文分词取长度不少于 2 的关键词。实体检索对前 5 个关键词分别执行 Cypher，查找名称包含关键词的节点和相连关系，去重并排除 Document 引用关系。子图查询从种子节点展开指定跳数。空间查询读取节点的 lng/lat 生成点，读取 path 生成折线。Neo4j 未连接或查询异常时返回空结构。", "KGPath[]、{nodes, edges}、{markers, polylines}。", "同一个 Neo4j 服务既服务在线问答，也服务独立图谱和地图页面。", "backend/src/services/kg.ts；backend/src/routes/kg.ts")
    add_module(doc, "3.6 大模型网关与提示词", "把检索得到的证据组织为模型能理解的上下文，并处理模型服务不可用的情况。", "问题、RAG 文本片段、KG 路径、检索模式。", "buildPrompt 按模式拼入文档片段（包含文档名、页码、相关度）和实体关系；SYSTEM_PROMPT 要求模型以地质问答助手身份作答。非流式模式优先调用 DeepSeek，遇到超时、连接、限流等可重试错误会尝试已配置的通义千问；全部失败则直接整理检索片段作为兜底答案。流式模式只调用第一个已注册模型，若失败则输出兜底答案，而不是跨模型续流。", "回答文本和来源列表，或一段段的流式文字。", "即使外部模型暂时不可用，系统仍能把检索到的原文呈现给用户。", "backend/src/services/llm.ts；backend/src/services/llm/gateway.ts；backend/src/services/llm/deepseek.ts；backend/src/services/llm/tongyi.ts")
    add_module(doc, "3.7 文档、日志与上传管理", "提供管理员查看资料状态、查看切片、管理实体、查看问答日志和上传 PDF 的能力。", "管理员请求；上传接口接收 file 字段的 PDF。", "SQLite 数据层读取 documents、chunks 和 qa_logs。上传路由用 multer 检查 PDF 类型与 50MB 限制，保存到 backend/uploads 后，以后台子进程启动 Python run_pipeline.py --file。管理路由对文档与实体提供列表、详情、更新和删除接口。", "统计数、文档/切片数据、日志、上传结果或更新结果。", "管理者可在页面上知道知识库是否已经产生切片，也能追踪系统使用情况。", "backend/src/routes/admin.ts；backend/src/routes/upload.ts；backend/src/db/sqlite.ts")

    doc.add_heading("四、数据端：把 PDF 变成可检索知识", level=1)
    doc.add_paragraph(
        "数据端使用 Python。它有两种工作状态：平时作为 Flask 检索服务运行；导入资料时运行若干离线脚本。"
        "两者都使用同一份 SQLite 文档库和 Neo4j 图谱，因此离线加工完成后，在线问答即可使用新资料。"
    )
    add_table(doc, ["数据模块", "输入", "处理", "输出"], [
        ["PDF 解析", "PDF 文件", "逐页提取并清理文本", "按页文本"],
        ["文本切片", "按页文本", "按自然边界切成重叠片段", "chunks"],
        ["索引和向量", "chunks", "TF-IDF 分词建模；BGE 编码", "SQLite、TF-IDF 模型、512 维向量"],
        ["实体关系抽取", "地质文本、地学词典、NER/LLM 配置", "词典、NER 和模型辅助抽取并去重过滤", "地质实体和关系"],
        ["图谱与本体", "实体、关系", "写入 Neo4j；构建 OWL 并推理", "节点、边、推理关系、OWL 文件"],
        ["Flask 检索服务", "问题、Top-K", "BGE 或 TF-IDF 相似度排序", "最相关文档片段"],
    ], [3.0, 3.0, 6.2, 6.3])

    add_module(doc, "4.1 PDF 解析", "把人可阅读的 PDF 转成程序可处理的逐页文本。", "PDF 路径。", "parse_pdf 使用 pdfplumber 按页提取文字，清理多余空白和由排版造成的断行，保留页码、标题和总页数。解析结果也可缓存为 parsed_texts.json，避免重复读取 PDF。", "{title, total_pages, pages:[{page_num, text}]}。", "后续切片、实体识别和定位来源页码都以它为基础。", "ml-service/scripts/parse_pdf.py")
    add_module(doc, "4.2 文本切片", "把长报告拆成适合检索和模型阅读的小段，同时尽量不断开一个完整语义。", "解析后的逐页文本；默认 chunk_size=512、overlap=128。", "切片器以约 512 字为目标，在窗口末尾优先寻找空行、句末标点、逗号等断点；相邻块保留约 128 字重叠。每块记录原文页码、序号、字符起止位置。", "带 doc_id、page、chunk_index、text、char_start、char_end 的片段列表。", "用户问到段落边界附近的内容时，重叠区能减少信息被切断的概率。", "ml-service/scripts/chunk_text.py")
    add_module(doc, "4.3 TF-IDF 索引、SQLite 与 BGE 语义向量", "为同一批切片建立两种检索能力：关键词匹配和语义相似度匹配。", "所有文档切片；增量模式时还会读取已有切片。", "build_index 用 jieba 分词，训练统一的 TF-IDF 词表，将 documents 和 chunks 写入 SQLite，并保存 tfidf_model.pkl。随后 embed_chunks 使用 bge-small-zh-v1.5 将每个文本块编码为归一化的 512 维 float32 向量，覆盖写入 chunks.vector，同时记录向量类型和维度。", "geo_knowledge.db、tfidf_model.pkl、每个 chunk 的 512 维向量。", "前者便于关键词检索，后者可找出说法不同但意思接近的材料。", "ml-service/scripts/build_index.py；ml-service/scripts/embed_chunks.py")
    add_module(doc, "4.4 地质实体与关系抽取", "从报告中识别矿产、岩石、构造、地质年代、矿床成因类型，并建立关系。", "解析后的地质文档、地学词典、NER 模型和 LLM 配置。", "主流程为地质文档调用 extract_hybrid：词典/正则先稳定识别已知术语，NER 用于补充识别，模型再结合文本识别实体间关系；结果会去重，并按允许的实体类型与关系类型过滤。用于训练 NER 的 dict_ner_geology 还会从 Neo4j 实体和内置词表生成 BIO 标注数据，再供训练和评估脚本使用。", "concepts（名称、类型、描述等）和 relations（起点、关系、终点）。", "把长文本里的“矿产—构造—岩石—年代”知识从段落形式转为可查询的结构。", "ml-service/scripts/extract_hybrid.py；ml-service/scripts/dict_ner_geology.py；ml-service/scripts/train_ner.py；ml-service/scripts/eval_ner.py")
    add_module(doc, "4.5 写入 Neo4j、构建 OWL 与推理", "把抽取结果存成图结构，并从显式关系推导可成立的隐含关系。", "文档、实体、关系，以及 Neo4j/OWL 配置。", "write_neo4j 按 Mineral、Rock、Structure、TimePeriod、DepositType 等标签 MERGE 节点，写入关系并建立 Document-REFERENCES-Entity 连接。build_ontology 将图谱内容转换为 OWL。run_reasoning 使用本体推理器得到新关系，并把推理关系标记为 inferred 后写回 Neo4j。", "Neo4j 节点和边、geo_planning.owl、带 inferred 标记的关系。", "图谱检索能直接得到关系链；图谱页面还可区分原文关系和推理补充关系。", "ml-service/scripts/write_neo4j.py；ml-service/scripts/build_ontology.py；ml-service/scripts/run_reasoning.py")
    add_module(doc, "4.6 Flask 运行时检索服务", "在用户提问时快速从 SQLite 中找出相关文本片段。", "POST /search 的 question 和 top_k。", "服务启动后按配置使用 BGE 或 TF-IDF。BGE 模式首次加载模型和 SQLite 中的全部 512 维向量；问题被编码为查询向量，与所有归一化文本向量做点积，排序后取前 K。TF-IDF 模式使用训练好的模型计算相似度。服务还提供 /health、/stats、/ner、/ontology/* 等开发与辅助接口。", "chunks 数组，元素含 id、text、doc_title、page、score；还可能给出低相关性提示。", "后端只需要一次 HTTP 调用即可获得可用于回答和引用的材料。", "ml-service/server.py")
    add_module(doc, "4.7 总控管线", "按正确顺序执行数据加工，避免人工逐个运行脚本。", "配置中的 PDF_FILES；可选 --file 单文件增量或 --from-step4。", "run_pipeline 依次执行：解析、切片、TF-IDF/SQLite、BGE 编码、实体关系抽取、写入 Neo4j、本体推理。--file 用于管理员上传后的增量处理；--from-step4 可从已有 SQLite 重建后半段图谱。", "终端统计和上述所有数据产物。", "新增资料后可以通过一条命令完成从 PDF 到知识库的转换。", "ml-service/scripts/run_pipeline.py")

    doc.add_heading("五、存储层：每种数据放在哪里", level=1)
    add_table(doc, ["位置", "存什么", "谁读写", "在问答中的作用"], [
        ["SQLite：ml-service/output/geo_knowledge.db", "documents、chunks、BGE 向量、qa_logs", "Python 管线写知识；Node 读文档/写日志", "提供可检索文本和来源页码，记录问答使用情况"],
        ["Neo4j", "矿产、岩石、构造、年代等节点与关系", "Python 管线写；Node KG 服务读写", "提供关系链、子图和可选空间属性"],
        ["模型/索引文件", "tfidf_model.pkl、NER 模型、OWL 文件", "Python 数据端", "支持检索、实体识别和本体推理"],
        ["浏览器 localStorage", "token、角色、前端会话消息", "前端", "维持登录态和本机对话历史"],
    ], [4.4, 5.1, 4.7, 4.3])
    doc.add_paragraph("需要注意：前端对话历史目前是浏览器本地保存，后端 /api/qa/history 暂时返回空数组；问答统计日志则保存于 SQLite。这两个用途不同，不能混为一谈。", style="Intense Quote")

    doc.add_heading("六、按代码顺序走一遍：一次提问怎样得到回答", level=1)
    doc.add_paragraph("以下以已登录用户输入“尾亚钒钛磁铁矿受什么构造控制？”并选择“混合检索”为例。具体命中的文档和实体会随当前数据库内容变化；这里展示的是代码固定执行的顺序。")
    add_number(doc, "用户在 QaView 的输入框发送问题。ChatInput 将文本交给 qa store；store 立即加入一条用户消息和一条空 AI 消息。")
    add_number(doc, "api/qa.ts 从 localStorage 取 JWT，向 POST /api/qa/ask 发送：{ question, stream: true, retrievalMode: 'hybrid' }。")
    add_number(doc, "后端 authMiddleware 验证 token。通过后，qa.ts 判断混合模式需要 RAG 与 KG 两条路径。")
    add_number(doc, "qa.ts 用 Promise.all 并行做三件事：调用 Flask /search 查文本片段；调用 Neo4j 的 searchEntities 查实体关系；调用 getSpatialResults 查坐标点和线。空间数据没有命中时，当前实现会按问题关键词尝试演示用 mock 数据。")
    add_number(doc, "Flask 在 BGE 模式下把问题编码为 512 维向量，与 SQLite 中所有 chunk 向量计算相似度，返回前 K 条，例如某一报告的某页片段及 score。KG 服务将问题分词，按“尾亚”“钒钛磁铁矿”“构造”等关键词在 Neo4j 查关系。")
    add_number(doc, "若文献和图谱都没有足够相关结果，后端直接流式返回一段引导用户提出地质问题的文字。若有结果，后端先发 meta 事件，其中有 RAG 条数、KG 条数和图谱路径。")
    add_number(doc, "llm.ts 把文献片段写进“文档检索上下文”，把关系写进“知识图谱上下文”，再加上原问题与回答要求，交给已配置的 DeepSeek。模型生成的每个文本片段由 onChunk 变成 SSE chunk 事件。")
    add_number(doc, "前端每接到一个 chunk，就追加到空 AI 消息，因此用户看到文字逐步出现。模型完成后，后端把 sources、kgContext、spatialData 放入 done 事件，并把本次命中数与耗时写入 qa_logs。")
    add_number(doc, "ChatMessage 渲染答案和来源卡片，KgMiniGraph 可展示关系；QaView 在有合法坐标时绘制地图标注。至此一轮问答结束。")
    doc.add_paragraph("这就是系统的实际分工：前端负责交互和显示，后端负责协调，数据端负责检索与知识加工，Neo4j/SQLite 负责保存可追溯的知识。", style="Intense Quote")

    doc.add_heading("七、接口和数据在各层之间怎样传递", level=1)
    add_table(doc, ["调用方向", "请求中的关键字段", "返回中的关键字段", "说明"], [
        ["前端 -> 后端登录", "username, password", "token, username, role", "获取登录身份"],
        ["前端 -> 后端问答", "question, retrievalMode, stream", "SSE: meta/chunk/done；或 JSON", "stream=true 时逐段接收"],
        ["后端 -> Flask 检索", "question, top_k", "chunks[{id,text,doc_title,page,score}]", "查本地文档片段"],
        ["后端 -> Neo4j", "关键词、跳数、实体 ID", "关系路径 / 子图 / 空间要素", "查实体关系和坐标"],
        ["后端 -> LLM", "系统提示词、上下文、问题", "回答文本或文字流", "组织语言，不替代本地检索"],
        ["数据管线 -> 存储", "PDF 文本、chunks、实体、关系", "SQLite/Neo4j/OWL 文件", "为在线问答预先准备数据"],
    ], [3.1, 5.3, 6.0, 4.1])

    doc.add_heading("八、实现边界和答辩时应如实说明的点", level=1)
    doc.add_paragraph("下面不是系统缺点清单，而是为了避免把“已经接通”和“预留能力”说成一件事。答辩时如实说明，反而更容易体现团队对系统边界的理解。")
    add_bullet(doc, "在线问答当前直接使用的是 Flask 的 /search 文献检索和 Node 端 Neo4j 图谱检索。NER 与 OWL 有脚本和 Flask 辅助接口，但不在每次 /api/qa/ask 的必经调用链上；它们主要用于离线建库和辅助验证。")
    add_bullet(doc, "独立的 /api/spatial/* GeoJSON 图层接口目前是预留骨架，尚未接入完整真实地质矢量图层。问答页的空间展示优先读取 Neo4j 坐标；没有真实坐标时，qa 路由会按关键词回退到演示 mock 数据。演示时应明确说明数据来源。")
    add_bullet(doc, "非流式大模型调用支持主模型失败后切到已配置备用模型；流式调用为保证事件流连续，目前只使用主模型，失败时改用检索原文的兜底回答。")
    add_bullet(doc, "上传 PDF 后会启动后台管线，接口先返回“已接收”。文档真正可检索需等待解析、切片、索引和图谱写入完成；控制台可通过 chunkCount 观察是否完成。")
    add_bullet(doc, "普通用户的聊天记录目前保存在浏览器 localStorage；SQLite 保存的是问答日志与统计，而不是完整服务端会话。")

    doc.add_heading("九、答辩时的简短讲法", level=1)
    doc.add_paragraph("如果老师让用一分钟概括系统，可以按下面顺序讲：")
    add_number(doc, "我们的目标不是让大模型凭空回答，而是让它基于地质报告和知识图谱回答。")
    add_number(doc, "数据端先把 PDF 解析、切片、向量化，形成 SQLite 文档库；同时抽取矿产、岩石、构造等实体关系，写入 Neo4j。")
    add_number(doc, "用户提问时，后端同时查文献和图谱，再把结果作为上下文交给大模型。")
    add_number(doc, "前端把答案流式显示，并把来源、关系和可用的空间信息展示出来。")
    add_number(doc, "因此每一层都有明确职责：前端负责交互，后端负责编排，数据端负责检索和建库，存储层负责记忆和追溯。")

    doc.add_heading("附录：主要源码对照", level=1)
    add_table(doc, ["类别", "重点文件", "说明"], [
        ["前端", "frontend/src/stores/qa.ts；frontend/src/api/qa.ts；frontend/src/views/QaView.vue", "问答消息状态、SSE 读取和页面展示"],
        ["前端", "frontend/src/views/GraphView.vue；frontend/src/views/MapView.vue；frontend/src/views/AdminView.vue", "图谱、地图和管理控制台"],
        ["后端", "backend/src/routes/qa.ts；backend/src/services/llm.ts；backend/src/services/llm/gateway.ts", "问答编排、提示词和模型调度"],
        ["后端", "backend/src/services/kg.ts；backend/src/services/tfidf.ts；backend/src/db/sqlite.ts", "Neo4j、Flask 检索调用和 SQLite 数据层"],
        ["后端", "backend/src/routes/auth.ts；backend/src/middleware/auth.ts；backend/src/routes/admin.ts；backend/src/routes/upload.ts", "鉴权、管理与上传"],
        ["数据端", "ml-service/server.py；ml-service/scripts/run_pipeline.py", "运行时检索服务与离线总控"],
        ["数据端", "ml-service/scripts/parse_pdf.py；chunk_text.py；build_index.py；embed_chunks.py", "PDF 到可检索向量的加工过程"],
        ["数据端", "ml-service/scripts/extract_hybrid.py；write_neo4j.py；build_ontology.py；run_reasoning.py", "实体关系、图谱和本体推理"],
    ], [3.0, 8.0, 7.5])

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
